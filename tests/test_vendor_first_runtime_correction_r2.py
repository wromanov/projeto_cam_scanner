from __future__ import annotations

from dataclasses import fields

import pytest

openpyxl = pytest.importorskip("openpyxl")

from cam_scanner.application.inventory_service import InventoryService
from cam_scanner.application.single_workflow import SingleWorkflow
from cam_scanner.cameras.preauth import ManufacturerResolver, PreAuthEvidence
from cam_scanner.cameras.registry import AdapterRegistry, build_default_registry
from cam_scanner.cameras.strategy import VendorFirstCollector
from cam_scanner.domain.enums import CollectionStatus
from cam_scanner.domain.models import (
    CameraResult,
    CameraTarget,
    CollectionAttempt,
    CollectionReport,
    EvidenceRecord,
)
from cam_scanner.excel.writer import ExcelWriter, InventoryOutputRow
from cam_scanner.terminal.ui import TerminalUI


def _complete_result(ip: str, manufacturer: str) -> CameraResult:
    return CameraResult(
        ip=ip,
        status=CollectionStatus.SUCCESS,
        manufacturer=manufacturer,
        model="Fixture model",
        serial="FIXTURE-001",
    )


def test_r2_single_hint_uses_eligible_vendor_before_any_onvif_discovery() -> None:
    events: list[str] = []
    target = CameraTarget("192.0.2.41", "fixture-user", "fixture-password")

    class ForbiddenDiscoverer:
        def discover_with_report(self, *_args, **_kwargs):
            raise AssertionError("eligible vendor hint must not require ONVIF preauth")

    class ForbiddenOnvif:
        def collect_report(self, *_args, **_kwargs):
            raise AssertionError("complete eligible vendor result needs no ONVIF call")

    from cam_scanner.cameras.manufacturers.hikvision import HikvisionAdapter

    adapter = HikvisionAdapter(session_factory=lambda: None)

    def fixture_collection(received: CameraTarget, *, context=None) -> CollectionReport:
        events.append("vendor")
        assert received == target
        assert context is not None
        report = CollectionReport(
            evidence=(
                EvidenceRecord("manufacturer", "Hikvision", "VENDOR", protocol="HTTPS"),
                EvidenceRecord("model", "Fixture model", "VENDOR", protocol="HTTPS"),
                EvidenceRecord("serial", "FIXTURE-001", "VENDOR", protocol="HTTPS"),
            ),
        )
        return report.with_attempt(
            CollectionAttempt(
                stage="HTTP",
                operation="GET /ISAPI/System/deviceInfo",
                protocol="HTTPS",
                outcome="SUCCESS",
            )
        )

    adapter.collect_report = fixture_collection  # type: ignore[method-assign]
    registry = AdapterRegistry()
    registry.register("Hikvision", adapter)
    collector = VendorFirstCollector(
        onvif_collector=ForbiddenOnvif(),
        registry=registry,
        discoverer=ForbiddenDiscoverer(),
    )
    service = InventoryService(collector)
    answers = iter([target.ip, target.username, "Hikvision", "", "3"])
    ui = TerminalUI(
        input_fn=lambda _prompt: next(answers),
        password_fn=lambda _prompt: target.password,
        output_fn=lambda _text: None,
    )

    SingleWorkflow(ui, service).run()

    assert events == ["vendor"]
    assert len(fields(_complete_result(target.ip, "Hikvision"))) == 14


def test_r2_strong_bosch_discovery_never_invokes_unsupported_adapter() -> None:
    events: list[str] = []
    target = CameraTarget("192.0.2.42", "fixture-user", "fixture-password")

    class EmptyDiscoverer:
        def discover_with_report(self, _target, *, context):
            events.append("discovery")
            return PreAuthEvidence(capabilities={"Manufacturer": "Bosch"}), CollectionReport().with_attempt(
                CollectionAttempt(
                    stage="DISCOVERY",
                    operation="GetCapabilities",
                    protocol="ONVIF",
                    outcome="SUCCESS",
                )
            )

    class GenericCollector:
        def collect_report(self, received, *, context):
            events.append("generic-onvif")
            context.auth.begin("ONVIF", "initial")
            return CollectionReport.from_camera_result(
                _complete_result(received.ip, "Generic Camera"),
                source="ONVIF_AUTHENTICATED",
            )

    registry = build_default_registry()
    bosch = registry.get("Bosch")
    assert bosch is not None

    def forbidden_bosch_collection(*_args, **_kwargs):
        events.append("bosch")
        raise AssertionError("registered Bosch stub is not eligible for collection")

    bosch.collect_report = forbidden_bosch_collection  # type: ignore[attr-defined]
    collector = VendorFirstCollector(
        onvif_collector=GenericCollector(),
        registry=registry,
        discoverer=EmptyDiscoverer(),
    )
    service = InventoryService(collector)
    context = service.create_context(target)
    try:
        report = collector.collect_report(target, context=context)
    finally:
        context.close()

    assert events == ["discovery", "generic-onvif"]
    assert sum(attempt.operation == "GetCapabilities" for attempt in report.attempts) == 1
    assert not any(attempt.operation == "vendor_adapter" for attempt in report.attempts)
    assert "manufacturer" in report.conflicts
    assert report.to_camera_result(target.ip).manufacturer is None


def test_r2_unknown_declared_hint_is_retained_without_selecting_a_vendor() -> None:
    resolution = ManufacturerResolver().resolve_with_evidence(
        "Unlisted model family", PreAuthEvidence()
    )

    assert resolution.declared == "Unlisted model family"
    assert resolution.selected is None
    assert resolution.resolved is None


def test_r2_single_unknown_hint_keeps_automatic_discovery_and_onvif_fallback() -> None:
    target = CameraTarget("192.0.2.57", "fixture-user", "fixture-password")
    events: list[str] = []

    class Discoverer:
        def discover_with_report(self, _target, *, context):
            events.append("discovery")
            return PreAuthEvidence(), CollectionReport().with_attempt(
                CollectionAttempt(
                    stage="DISCOVERY",
                    operation="GetCapabilities",
                    protocol="ONVIF",
                    outcome="SUCCESS",
                )
            )

    class GenericOnvif:
        def collect_report(self, received, *, context):
            events.append("onvif")
            return CollectionReport.from_camera_result(
                _complete_result(received.ip, "Hikvision"),
                source="ONVIF_AUTHENTICATED",
            )

    collector = VendorFirstCollector(
        onvif_collector=GenericOnvif(),
        discoverer=Discoverer(),
    )
    context = InventoryService(collector).create_context(
        target, manufacturer_hint="Unlisted model family"
    )
    try:
        report = collector.collect_report(target, context=context)
    finally:
        context.close()

    assert events == ["discovery", "onvif"]
    assert report.manufacturer_resolution is not None
    assert report.manufacturer_resolution.declared == "Unlisted model family"
    assert report.manufacturer_resolution.observed == "Hikvision"
    assert report.to_camera_result(target.ip).manufacturer == "Hikvision"


@pytest.mark.parametrize(
    ("raw_hint", "declared", "selected"),
    [("HIK", "HIK", "Hikvision"), ("Unlisted family", "Unlisted family", None)],
)
def test_r2_multi_reader_preserves_literal_alias_and_unknown_hints(
    tmp_path, raw_hint: str, declared: str, selected: str | None
) -> None:
    from openpyxl import Workbook

    from cam_scanner.excel.reader import ExcelReader

    source = tmp_path / "multi-hints.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(("IP", "USERNAME", "PASSWORD", "FABRICANTE"))
    worksheet.append(("192.0.2.55", "fixture-user", "fixture-password", raw_hint))
    workbook.save(source)
    workbook.close()

    row = next(ExcelReader().read(source).rows)
    resolution = ManufacturerResolver().resolve_with_evidence(
        row.manufacturer_hint, PreAuthEvidence()
    )

    assert row.manufacturer_hint == raw_hint
    assert resolution.declared == declared
    assert resolution.selected == selected


def test_r2_non_tls_vendor_failure_preserves_attempt_and_uses_allowed_onvif_fallback() -> None:
    target = CameraTarget("192.0.2.47", "fixture-user", "fixture-password")
    events: list[str] = []

    class ForbiddenDiscoverer:
        def discover_with_report(self, *_args, **_kwargs):
            raise AssertionError("eligible hint starts directly with its adapter")

    class GenericCollector:
        def collect_report(self, received, *, context):
            events.append("onvif")
            return CollectionReport.from_camera_result(
                _complete_result(received.ip, "Hikvision"),
                source="ONVIF_AUTHENTICATED",
            )

    from cam_scanner.cameras.manufacturers.hikvision import HikvisionAdapter

    adapter = HikvisionAdapter(session_factory=lambda: None)

    def failed_vendor(_target, *, context=None) -> CollectionReport:
        events.append("vendor")
        return CollectionReport().with_attempt(
            CollectionAttempt(
                stage="CONNECT",
                operation="vendor_adapter",
                outcome="FAILED",
                error_category="CONNECTION_REFUSED",
            )
        )

    adapter.collect_report = failed_vendor  # type: ignore[method-assign]
    registry = AdapterRegistry()
    registry.register("Hikvision", adapter)
    collector = VendorFirstCollector(
        onvif_collector=GenericCollector(),
        registry=registry,
        discoverer=ForbiddenDiscoverer(),
    )
    service = InventoryService(collector)
    context = service.create_context(target, manufacturer_hint="Hikvision")
    try:
        report = collector.collect_report(target, context=context)
    finally:
        context.close()

    assert events == ["vendor", "onvif"]
    assert any(
        attempt.error_category == "CONNECTION_REFUSED" for attempt in report.attempts
    )
    assert report.to_camera_result(target.ip).status is CollectionStatus.SUCCESS


@pytest.mark.parametrize("tls_source", ["discovery", "vendor"])
def test_r2_tls_error_is_terminal_for_all_later_network_attempts(tls_source: str) -> None:
    target = CameraTarget("192.0.2.48", "fixture-user", "fixture-password")
    events: list[str] = []

    class TLSDiscoverer:
        def discover_with_report(self, *_args, **_kwargs):
            events.append("discovery")
            report = CollectionReport().with_attempt(
                CollectionAttempt(
                    stage="CONNECT",
                    operation="preauth_discovery",
                    protocol="ONVIF",
                    outcome="FAILED",
                    error_category="TLS_ERROR",
                )
            )
            return PreAuthEvidence(), report

    class ForbiddenOnvif:
        def collect_report(self, *_args, **_kwargs):
            events.append("onvif")
            raise AssertionError("no fallback is allowed after TLS_ERROR")

    from cam_scanner.cameras.manufacturers.hikvision import HikvisionAdapter

    adapter = HikvisionAdapter(session_factory=lambda: None)

    def tls_vendor(_target, *, context=None) -> CollectionReport:
        events.append("vendor")
        return CollectionReport().with_attempt(
            CollectionAttempt(
                stage="CONNECT",
                operation="vendor_adapter",
                protocol="HTTPS",
                outcome="FAILED",
                error_category="TLS_ERROR",
            )
        )

    adapter.collect_report = tls_vendor  # type: ignore[method-assign]
    registry = AdapterRegistry()
    registry.register("Hikvision", adapter)
    collector = VendorFirstCollector(
        onvif_collector=ForbiddenOnvif(),
        registry=registry,
        discoverer=TLSDiscoverer(),
    )
    service = InventoryService(collector)
    hint = "Hikvision" if tls_source == "vendor" else None
    context = service.create_context(target, manufacturer_hint=hint)
    try:
        report = collector.collect_report(target, context=context)
        assert context.tls_failed
    finally:
        context.close()

    expected = ["vendor"] if tls_source == "vendor" else ["discovery"]
    assert events == expected
    assert any(attempt.error_category == "TLS_ERROR" for attempt in report.attempts)


def test_r2_vendor_observation_preserves_a_strong_hint_mismatch() -> None:
    target = CameraTarget("192.0.2.49", "fixture-user", "fixture-password")
    from cam_scanner.cameras.manufacturers.hikvision import HikvisionAdapter

    adapter = HikvisionAdapter(session_factory=lambda: None)

    def vendor_reports_other_identity(_target, *, context=None) -> CollectionReport:
        return CollectionReport(
            evidence=(
                EvidenceRecord("manufacturer", "Dahua", "VENDOR", protocol="HTTPS"),
                EvidenceRecord("model", "Observed model", "VENDOR", protocol="HTTPS"),
                EvidenceRecord("serial", "OBSERVED-001", "VENDOR", protocol="HTTPS"),
            )
        )

    adapter.collect_report = vendor_reports_other_identity  # type: ignore[method-assign]
    registry = AdapterRegistry()
    registry.register("Hikvision", adapter)
    collector = VendorFirstCollector(registry=registry)
    service = InventoryService(collector)
    context = service.create_context(target, manufacturer_hint="Hikvision")
    try:
        report = collector.collect_report(target, context=context)
    finally:
        context.close()

    resolution = report.manufacturer_resolution
    assert resolution is not None
    assert resolution.declared == "Hikvision"
    assert resolution.observed == "Dahua"
    assert resolution.resolved == "Dahua"
    assert resolution.mismatch is True


def test_r2_conflicting_strong_discovery_is_independent_of_response_order() -> None:
    target = CameraTarget("192.0.2.50", "fixture-user", "fixture-password")

    class GenericCollector:
        def __init__(self) -> None:
            self.calls = 0

        def collect_report(self, received, *, context=None) -> CollectionReport:
            self.calls += 1
            return CollectionReport.from_camera_result(
                _complete_result(received.ip, "Generic Camera"),
                source="ONVIF_AUTHENTICATED",
            )

    resolutions = []
    for families in (("Axis", "Dahua"), ("Dahua", "Axis")):
        class OrderedDiscoverer:
            def discover(self, _target, *, families=families):
                return PreAuthEvidence(
                    capabilities=[{"manufacturer": family} for family in families]
                )

        generic = GenericCollector()
        registry = AdapterRegistry()
        registry.register("Axis", object())
        registry.register("Dahua", object())
        report = VendorFirstCollector(
            onvif_collector=generic,
            registry=registry,
            discoverer=OrderedDiscoverer(),
        ).collect_report(target)
        resolution = report.manufacturer_resolution
        assert resolution is not None
        assert resolution.selected is None
        assert resolution.confidence == "CONFLICTING"
        assert report.to_camera_result(target.ip).manufacturer is None
        assert generic.calls == 1
        resolutions.append(resolution)

    assert resolutions[0] == resolutions[1]


def test_r2_authenticated_onvif_evidence_reconciles_final_manufacturer_resolution() -> None:
    target = CameraTarget("192.0.2.54", "fixture-user", "fixture-password")

    class EmptyDiscoverer:
        def discover_with_report(self, _target, *, context):
            return PreAuthEvidence(), CollectionReport()

    class IdentifyingOnvif:
        def collect_report(self, received, *, context=None):
            return CollectionReport.from_camera_result(
                _complete_result(received.ip, "Hikvision"),
                source="ONVIF_AUTHENTICATED",
            )

    report = VendorFirstCollector(
        onvif_collector=IdentifyingOnvif(),
        discoverer=EmptyDiscoverer(),
    ).collect_report(target)

    assert report.to_camera_result(target.ip).manufacturer == "Hikvision"
    assert report.manufacturer_resolution is not None
    assert report.manufacturer_resolution.observed == "Hikvision"
    assert report.manufacturer_resolution.resolved == "Hikvision"


def test_r2_single_diagnostics_display_declared_observed_and_resolved_mismatch() -> None:
    output: list[str] = []
    resolution = ManufacturerResolver().resolve_with_evidence(
        "Hikvision", PreAuthEvidence(capabilities={"Manufacturer": "Axis"})
    )
    report = CollectionReport().with_manufacturer_resolution(resolution)
    ui = TerminalUI(input_fn=lambda _prompt: "", output_fn=output.append)

    ui.render_report(report, ip="192.0.2.56", duration=0.1)

    rendered = "\n".join(output)
    assert "Fabricante declarado: Hikvision" in rendered
    assert "Fabricante observado: Axis" in rendered
    assert "Fabricante resolvido: Axis" in rendered
    assert "Divergência entre fabricante declarado e observado." in rendered


def test_r2_weak_manufacturer_signals_do_not_route_vendor_without_a_hint() -> None:
    resolver = ManufacturerResolver()
    resolution = resolver.resolve_with_evidence(
        None,
        PreAuthEvidence(
            capabilities={"Server": "Hikvision", "realm": "Axis", "deviceName": "Dahua"}
        ),
    )

    assert resolution.selected is None
    assert resolution.resolved is None
    assert resolution.confidence == "WEAK"


def test_r2_diagnostics_sheet_is_ordered_sanitized_and_formula_safe(tmp_path) -> None:
    secret = "fixture-secret-token"

    def unsafe_attempt() -> CollectionAttempt:
        attempt = CollectionAttempt(
            stage="HTTP",
            operation="collect",
            protocol="HTTPS",
            outcome="FAILED",
            error_category="TLS_ERROR",
            protocol_code="NotAuthorized",
            http_status=401,
            duration=0.25,
        )
        # Simulate corrupted legacy in-memory attempts at the XLSX boundary.
        object.__setattr__(attempt, "stage", "=1+1")
        object.__setattr__(attempt, "operation", "=HYPERLINK(\"https://bad\")")
        object.__setattr__(attempt, "error_category", "@secret")
        object.__setattr__(attempt, "protocol_code", secret)
        return attempt

    unsafe_attempts = [unsafe_attempt()]
    for marker in ("+1", "-1", "@SUM(A1)", "\t=1+1", "\r=1+1", "\x01=1"):
        attempt = CollectionAttempt(
            stage="HTTP",
            operation="collect",
            protocol="HTTPS",
            outcome="FAILED",
            error_category="TLS_ERROR",
        )
        object.__setattr__(attempt, "operation", marker)
        object.__setattr__(attempt, "protocol_code", marker)
        unsafe_attempts.append(attempt)

    earlier = CollectionReport(attempts=tuple(unsafe_attempts)).with_attempt(
        CollectionAttempt(
            stage="DISCOVERY",
            operation="GetCapabilities",
            protocol="ONVIF",
            outcome="SUCCESS",
            duration=0.1,
        )
    ).with_attempt(
        CollectionAttempt(
            stage="CONNECT",
            operation="preauth_discovery",
            protocol="ONVIF",
            outcome="FAILED",
            error_category="TLS_ERROR",
        )
    ).with_attempt(
        CollectionAttempt(
            stage="CONNECT",
            operation="vendor_adapter",
            protocol="HTTPS",
            outcome="FAILED",
            error_category="TLS_ERROR",
        )
    )
    later = CollectionReport().with_attempt(
        CollectionAttempt(
            stage="CONNECT",
            operation="create_onvif_client",
            protocol="ONVIF",
            outcome="FAILED",
            error_category="CONNECTION_REFUSED",
        )
    )
    earlier = earlier.with_manufacturer_resolution(
        ManufacturerResolver().resolve_with_evidence(
            "Hikvision",
            PreAuthEvidence(capabilities=[{"manufacturer": "Axis"}]),
        )
    )
    destination = tmp_path / "vendor-first-diagnostics.xlsx"

    ExcelWriter().write(
        destination,
        [
            InventoryOutputRow(6, _complete_result("192.0.2.46", "Hikvision"), report=later),
            InventoryOutputRow(3, _complete_result("192.0.2.43", "Axis"), report=earlier),
        ],
    )

    workbook = openpyxl.load_workbook(destination, data_only=False)
    try:
        diagnostics = workbook["DIAGNOSTICOS"]
        inventory = workbook["INVENTARIO"]
        assert tuple(cell.value for cell in diagnostics[1]) == (
            "LINHA",
            "ETAPA",
            "PROTOCOLO",
            "OPERACAO",
            "OUTCOME",
            "CATEGORIA",
            "DURACAO",
            "CODIGO",
            "STATUS",
        )
        assert [diagnostics.cell(row, 1).value for row in range(2, 13)] == [3] * 10 + [6]
        assert diagnostics[2][1].value == "N/A"
        assert all(diagnostics.cell(row, 2).value == "HTTP" for row in range(3, 9))
        assert all(diagnostics.cell(row, 4).value == "N/A" for row in range(2, 9))
        assert diagnostics[2][5].value == "N/A"
        assert all(diagnostics.cell(row, 6).value == "TLS_ERROR" for row in range(3, 9))
        assert all(diagnostics.cell(row, 8).value == "N/A" for row in range(2, 9))
        assert diagnostics[2][8].value == 401
        assert diagnostics[10][3].value == "preauth_discovery"
        assert diagnostics[11][3].value == "vendor_adapter"
        assert "diverge do observado" in inventory[3][13].value
        assert "Hikvision" not in inventory[3][13].value
        assert "Axis" not in inventory[3][13].value
        all_values = [cell.value for row in diagnostics.iter_rows() for cell in row]
        assert secret not in str(all_values)
        assert "HYPERLINK" not in str(all_values)
        assert all(cell.data_type != "f" for row in diagnostics.iter_rows() for cell in row)
    finally:
        workbook.close()

