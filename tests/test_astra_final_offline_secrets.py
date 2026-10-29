"""Offline output-boundary evidence for the Astra secret canaries.

The only network fixture binds to loopback and keeps certificate verification
enabled.  Synthetic tokens are checked without printing the matching content.
"""
from __future__ import annotations

from dataclasses import fields
from pathlib import Path
from zipfile import ZipFile

import openpyxl
import pytest

from cam_scanner.application.inventory_service import InventoryService
from cam_scanner.application.multi_workflow import MultiWorkflow
from cam_scanner.application.single_workflow import SingleWorkflow
from cam_scanner.cameras.transport import TargetTransportPolicy
from cam_scanner.config.loader import load_settings
from cam_scanner.domain.enums import CollectionStatus, MainMenuChoice
from cam_scanner.domain.errors import ConfigurationError
from cam_scanner.domain.models import CameraResult, CameraTarget
from cam_scanner.logging.setup import close_logging, configure_logging
from cam_scanner.terminal.ui import TerminalUI
from tests.test_astra_gpn_joint_integration import (
    _CapturingInventory,
    _HttpsFixture,
    _inventory,
    _NoMedia,
    _Progress,
    _trust_fixture_ca,
)

_USER = "fixture-account-a91"
_PASSWORD = "fixture-password-b92"
_HEADER = "fixture-header-c93"
_BODY = "fixture-body-d94"
_NONCE = "fixture-nonce-e95"
_SECRET = "fixture-secret-f96"
_URL = "https://fixture.invalid/private-g97"
_FORMULA = "fixture-formula-h98"
_CANARIES = (_USER, _PASSWORD, _HEADER, _BODY, _NONCE, _SECRET, _URL, _FORMULA)


def _assert_clean(value: str | bytes, surface: str = "output") -> None:
    """Never show the scanned content or the particular matching canary on failure."""
    payload = value if isinstance(value, bytes) else value.encode("utf-8")
    if any(canary.encode("utf-8") in payload for canary in _CANARIES):
        pytest.fail(f"A sensitive synthetic marker reached {surface}.", pytrace=False)


def _assert_clean_workbook(path: Path) -> None:
    with ZipFile(path) as archive:
        for entry in archive.infolist():
            _assert_clean(archive.read(entry.filename), "XLSX archive")
    workbook = openpyxl.load_workbook(path, data_only=False)
    try:
        assert set(workbook.sheetnames) == {"INVENTARIO", "RESUMO", "DIAGNOSTICOS"}
        assert len(next(iter(workbook["DIAGNOSTICOS"].iter_rows(values_only=True)))) == 9
        assert workbook["DIAGNOSTICOS"].max_row > 1
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    assert cell.data_type != "f"
                    if cell.value is not None:
                        _assert_clean(str(cell.value), "XLSX cell")
    finally:
        workbook.close()


def _terminal(inputs: list[str], outputs: list[str]) -> TerminalUI:
    answers = iter(inputs)
    return TerminalUI(
        input_fn=lambda _prompt: next(answers),
        password_fn=lambda _prompt: _PASSWORD,
        output_fn=outputs.append,
    )


def _hostile_cgi(handler) -> None:
    if not handler.headers.get("Authorization"):
        handler.send_response(401)
        handler.send_header(
            "WWW-Authenticate", f'Digest realm="loopback", nonce="{_NONCE}", qop="auth"',
        )
        handler.send_header("Content-Length", "0")
        handler.end_headers()
        return
    body = (
        "NAME=WV-S2531L\nSERIAL=SAFE-SERIAL\nVERSION=9.99\n"
        f"USERNAME={_USER}\nPASSWORD={_PASSWORD}\nAUTHORIZATION={_HEADER}\n"
        f"REQUEST_BODY={_BODY}\nNONCE={_NONCE}\nSECRET={_SECRET}\nWEBURL={_URL}\n"
    ).encode()
    handler.send_response(200)
    handler.send_header("Content-Type", "text/plain")
    handler.send_header("Content-Length", str(len(body)))
    handler.end_headers()
    handler.wfile.write(body)


def test_single_multi_hostile_response_never_reaches_outputs(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path,
) -> None:
    """Real workflows, terminal, transport, adapters, log and XLSX share canaries."""
    log_dir = tmp_path / "logs"
    logger = configure_logging(log_directory=log_dir)
    single_output: list[str] = []
    multi_output: list[str] = []
    source, destination = tmp_path / "input.xlsx", tmp_path / "output.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(("IP", "USERNAME", "PASSWORD", "FABRICANTE", "MODELO"))
    sheet.append(("127.0.0.1", _USER, _PASSWORD, "Panasonic", "WV-S2531L"))
    sheet.append(("127.0.0.1", _USER, _PASSWORD, "Panasonic", f'= "{_FORMULA}"'))
    workbook.save(source)
    workbook.close()

    try:
        server = _HttpsFixture("digest")
        monkeypatch.setattr(server.httpd.RequestHandlerClass, "_cgi", _hostile_cgi)
        with server:
            _trust_fixture_ca(monkeypatch, server.port)
            single_inventory = _CapturingInventory(_inventory(server.port))
            single_inventory.service._logger = logger
            single_ui = _terminal(
                ["127.0.0.1", _USER, "Panasonic", "WV-S2531L", "3"], single_output,
            )
            assert SingleWorkflow(single_ui, single_inventory, _NoMedia()).run() is MainMenuChoice.EXIT

            multi_inventory = _CapturingInventory(_inventory(server.port))
            multi_inventory.service._logger = logger
            multi_ui = _terminal([str(source), str(destination)], multi_output)
            workflow = MultiWorkflow(
                multi_ui, multi_inventory, _NoMedia(), progress_reporter=_Progress(), max_workers=1,
            )
            workflow.run()

        cgi = [item for item in server.requests if str(item["path"]).startswith("/cgi-bin/")]
        assert len(cgi) == 4
        assert sum(item["authorization"] is None for item in cgi) == 2
        assert all(item["scheme"] == "https" for item in server.requests)
        assert len(workflow.completed_rows) == 2
        assert workflow.completed_rows[0].result.status is CollectionStatus.PARTIAL_SUCCESS
        assert any(
            attempt.operation == "validate_model_hint" and attempt.error_category == "INVALID_INPUT"
            for attempt in workflow.completed_rows[1].report.attempts
        )
        assert len(fields(CameraResult)) == 14
        assert destination.exists()
        for output in (single_output, multi_output):
            _assert_clean("\n".join(output), "terminal")
        for context in (*single_inventory.contexts, *multi_inventory.contexts):
            _assert_clean(repr(context), "context repr")
            _assert_clean(repr(context.target), "target repr")
            _assert_clean(repr(context.report), "report repr")
            _assert_clean(repr(context.report.to_camera_result("127.0.0.1")), "result repr")
        for row in workflow.completed_rows:
            _assert_clean(repr(row), "output row repr")
        _assert_clean_workbook(destination)
    finally:
        close_logging()
    for log_file in log_dir.glob("cam_scanner_*.log*"):
        _assert_clean(log_file.read_bytes(), "log")


class _UnexpectedCollector:
    def collect_report(self, target: CameraTarget, *, context=None):
        del target, context
        payload = (
            f"username={_USER}\npassword={_PASSWORD}\nAuthorization: Bearer {_HEADER}\n"
            f"url={_URL}\nnonce={_NONCE}\nbody={_BODY}\nsecret={_SECRET}"
        )
        raise RuntimeError(payload)


@pytest.mark.parametrize("mode", ["single", "multi"])
def test_unexpected_collector_error_keeps_raw_payload_out_of_log_and_outputs(
    tmp_path: Path, mode: str,
) -> None:
    """An arbitrary upstream error is checked at real log and workflow boundaries."""
    log_dir = tmp_path / "logs"
    logger = configure_logging(log_directory=log_dir)
    output: list[str] = []
    try:
        service = InventoryService(_UnexpectedCollector(), logger=logger)
        if mode == "single":
            ui = _terminal(["127.0.0.1", _USER, "", "", "3"], output)
            assert SingleWorkflow(ui, service).run() is MainMenuChoice.EXIT
        else:
            source, destination = tmp_path / "exception-input.xlsx", tmp_path / "exception-output.xlsx"
            workbook = openpyxl.Workbook()
            workbook.active.append(("IP", "USERNAME", "PASSWORD"))
            workbook.active.append(("127.0.0.1", _USER, _PASSWORD))
            workbook.save(source)
            workbook.close()
            ui = _terminal([str(source), str(destination)], output)
            workflow = MultiWorkflow(ui, service, _NoMedia(), progress_reporter=_Progress(), max_workers=1)
            workflow.run()
            assert len(workflow.completed_rows) == 1
            assert workflow.completed_rows[0].result.status is CollectionStatus.FAILED
            _assert_clean(repr(workflow.completed_rows), "output repr")
            _assert_clean_workbook(destination)
    finally:
        close_logging()
    assert output
    _assert_clean("\n".join(output), "terminal")
    files = tuple(log_dir.glob("cam_scanner_*.log*"))
    assert files
    for log_file in files:
        assert b"UNEXPECTED" in log_file.read_bytes()
        _assert_clean(log_file.read_bytes(), "log")


@pytest.mark.parametrize("key", ["username", "password", "authorization", "nonce", "realm", "body", "cookie", "secret"])
def test_profile_rejects_secret_fields_without_leaking_error_or_starting_network(
    key: str, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    from cam_scanner.cameras import transport as transport_module

    config = tmp_path / "profile.toml"
    config.write_text(
        '[transport]\n[[transport.profiles]]\nid="fixture-policy"\n'
        'targets=["127.0.0.1"]\ncomponent="vendor_panasonic"\nprotocol="http"\n'
        'port=8080\nauth_mode="NATIVE_DIGEST_EXISTING"\n'
        'authorization_ref="fixture-approval"\nhttp_risk_accepted=true\n'
        f'{key}="{_SECRET}"\n', encoding="utf-8",
    )

    def network_constructor(*_args, **_kwargs):
        pytest.fail("A rejected profile created a network-capable object.", pytrace=False)

    monkeypatch.setattr(transport_module.requests, "Session", network_constructor)
    monkeypatch.setattr(TargetTransportPolicy, "__init__", network_constructor)
    with pytest.raises(ConfigurationError) as caught:
        load_settings(config)
    _assert_clean(str(caught.value))
    _assert_clean(repr(caught.value))
    assert caught.value.__suppress_context__
