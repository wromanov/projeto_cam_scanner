"""Synthetic manufacturer routing through the real SINGLE CLI vertical slice."""
from __future__ import annotations

import importlib
import logging

import pytest
from openpyxl import Workbook, load_workbook

from cam_scanner.domain.enums import CollectionStatus

_FAMILIES = (
    ("Hikvision", "hikCapabilities"),
    ("Axis", "axisCapabilities"),
    ("Dahua", "dahuaCapabilities"),
    ("Panasonic", "panasonicCapabilities"),
    ("Samsung/Hanwha", "hanwhaCapabilities"),
    ("Bosch", "boschCapabilities"),
)
_VENDOR_COMPONENTS = {
    "Hikvision": "vendor_hikvision",
    "Axis": "vendor_axis",
    "Dahua": "vendor_dahua",
    "Panasonic": "vendor_panasonic",
    "Samsung/Hanwha": "vendor_hanwha",
}


class _SyntheticTransport:
    """Keep the product adapters real while terminating every network request locally."""

    def __init__(self, context, requests_seen: list[tuple[str, str, str]]) -> None:
        self.context = context
        self.requests_seen = requests_seen

    def endpoint(self, component: str, *, protocol: str = "https", path: str = "/") -> str:
        scheme = "http" if protocol == "http" else "https"
        port = ":80" if scheme == "http" else ""
        return f"{scheme}://{self.context.target.ip}{port}{path}"

    def authorize(self, component: str, url: str) -> str:
        return url

    def request(self, component: str, method: str, url: str, *, source: str, **_options):
        self.context.check(source)
        self.requests_seen.append((self.context.target.ip, component, method))
        # A local 404 exercises the adapters' documented unsupported-endpoint fallback.
        return _SyntheticResponse(404)


class _SyntheticResponse:
    def __init__(self, status_code: int) -> None:
        self.status_code = status_code
        self.headers: dict[str, str] = {}

    def close(self) -> None:
        pass


class _SyntheticControlledOnvifClient:
    """Stand-in at the controlled ONVIF boundary, with no socket-capable client."""

    def __init__(self, context, *, authenticated: bool, calls: list[tuple[str, str, bool]]) -> None:
        self.context = context
        self.authenticated = authenticated
        self.calls = calls

    def devicemgmt(self):
        return _SyntheticDeviceService(self.context, self.authenticated, self.calls)


class _SyntheticDeviceService:
    def __init__(self, context, authenticated: bool, calls: list[tuple[str, str, bool]]) -> None:
        self.context = context
        self.authenticated = authenticated
        self.calls = calls

    def GetCapabilities(self, **_kwargs: object) -> dict[str, object]:
        self.calls.append((self.context.target.ip, "GetCapabilities", self.authenticated))
        family = _family_for_ip(self.context.target.ip)
        return {"Device": {"Manufacturer": family}}

    def GetDeviceInformation(self) -> dict[str, str]:
        self.calls.append((self.context.target.ip, "GetDeviceInformation", self.authenticated))
        assert self.authenticated
        self.context.auth.observe_http(200, "ONVIF", authenticated=True)
        family = _family_for_ip(self.context.target.ip)
        last_octet = self.context.target.ip.rsplit(".", 1)[-1]
        return {
            "Manufacturer": family,
            "Model": f"Synthetic {family} model",
            "SerialNumber": f"serial-{last_octet}",
            "FirmwareVersion": "synthetic-1.0",
        }


def _family_for_ip(ip: str) -> str:
    index = int(ip.rsplit(".", 1)[-1]) - 10
    return _FAMILIES[index][0]


def _install_synthetic_boundaries(monkeypatch, requests_seen, calls) -> None:
    from cam_scanner.application.inventory_service import InventoryService
    from cam_scanner.snapshot.service import SnapshotCaptureError, SnapshotFailure, SnapshotService

    def transport_factory(context):
        return _SyntheticTransport(context, requests_seen)

    monkeypatch.setattr(InventoryService, "create_context", lambda self, target, **_kwargs: __import__(
        "cam_scanner.application.collection_context", fromlist=["CollectionContext"]
    ).CollectionContext(target, self._config, transport_factory=transport_factory))
    factory = lambda context, *, authenticated: _SyntheticControlledOnvifClient(
        context, authenticated=authenticated, calls=calls
    )
    # Both call sites resolve this symbol at request time; no third-party defaults are patched.
    monkeypatch.setattr("cam_scanner.cameras.preauth.ControlledOnvifClient", factory)
    monkeypatch.setattr("cam_scanner.cameras.onvif.adapter.ControlledOnvifClient", factory)

    class UnsupportedProvider:
        def capture(self, _target, **_kwargs):
            raise SnapshotCaptureError(SnapshotFailure.UNSUPPORTED)

    main_module = importlib.import_module("cam_scanner.main")
    monkeypatch.setattr(
        main_module,
        "SnapshotService",
        lambda: SnapshotService(UnsupportedProvider(), UnsupportedProvider(), UnsupportedProvider()),
    )
    monkeypatch.setattr(
        main_module,
        "configure_logging",
        lambda **_kwargs: logging.getLogger("synthetic-pre-p12-matrix"),
    )


@pytest.mark.parametrize("family, _capability", _FAMILIES)
def test_single_entrypoint_discovers_routes_falls_back_and_renders_identity(
    family: str,
    _capability: str,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    """The family comes from anonymous ONVIF evidence; registry and adapters remain real."""
    calls: list[tuple[str, str, bool]] = []
    requests_seen: list[tuple[str, str, str]] = []
    _install_synthetic_boundaries(monkeypatch, requests_seen, calls)

    index = next(i for i, item in enumerate(_FAMILIES) if item[0] == family)
    ip = f"192.0.2.{10 + index}"
    password = f"synthetic-secret-{index}-do-not-print"
    answers = iter(["1", ip, "synthetic-user", "", "", "3"])
    monkeypatch.setattr("builtins.input", lambda _prompt: next(answers))
    monkeypatch.setattr("getpass.getpass", lambda _prompt: password)

    main_module = importlib.import_module("cam_scanner.main")
    assert main_module.main() == 0

    terminal = capsys.readouterr().out
    target_calls = [operation for host, operation, _authenticated in calls if host == ip]
    assert target_calls.count("GetCapabilities") == 1
    assert target_calls.count("GetDeviceInformation") == 1
    capabilities_call = next(item for item in calls if item[0] == ip and item[1] == "GetCapabilities")
    information_call = next(
        item for item in calls if item[0] == ip and item[1] == "GetDeviceInformation"
    )
    assert capabilities_call[2] is False
    assert information_call[2] is True
    assert f"Fabricante: {family}" in terminal
    assert f"Modelo: Synthetic {family} model" in terminal
    assert f"Número de série: serial-{ip.rsplit('.', 1)[-1]}" in terminal
    assert password not in terminal
    assert "Traceback" not in terminal

    vendor_requests = [request for request in requests_seen if request[0] == ip]
    if family in {"Bosch", "Panasonic"}:
        assert vendor_requests == []
    else:
        assert len(vendor_requests) == 1
        assert vendor_requests[0][1] == _VENDOR_COMPONENTS[family]
        assert vendor_requests[0][2] in {"GET", "POST"}

    # The tested identity is produced by the real collector/report and stays typed end to end.
    from cam_scanner.domain.models import CameraResult

    assert len(CameraResult.__dataclass_fields__) == 14
    assert CollectionStatus.SUCCESS.value in terminal or "Status: PARTIAL_SUCCESS" in terminal


def test_multi_entrypoint_exports_synthetic_family_results_in_input_order(
    tmp_path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    """Exercise the real XLSX reader/writer without putting credentials in output."""
    calls: list[tuple[str, str, bool]] = []
    requests_seen: list[tuple[str, str, str]] = []
    _install_synthetic_boundaries(monkeypatch, requests_seen, calls)

    source = tmp_path / "synthetic-cameras.xlsx"
    destination = tmp_path / "synthetic-inventory.xlsx"
    input_book = Workbook()
    input_sheet = input_book.active
    input_sheet.append(("IP", "USERNAME", "PASSWORD"))
    ordered_ips = [f"192.0.2.{10 + index}" for index in range(len(_FAMILIES))]
    secrets = [f"synthetic-secret-{index}-do-not-export" for index in range(len(_FAMILIES))]
    for ip, password in zip(ordered_ips, secrets, strict=True):
        input_sheet.append((ip, "synthetic-user", password))
    input_book.save(source)
    input_book.close()

    answers = iter(["2", str(source), str(destination), "3"])
    monkeypatch.setattr("builtins.input", lambda _prompt: next(answers))
    main_module = importlib.import_module("cam_scanner.main")

    assert main_module.main() == 0

    terminal = capsys.readouterr().out
    assert "MULTI concluído: 6 linhas processadas; 6 concluídas, 0 falhas." in terminal
    assert all(password not in terminal for password in secrets)
    assert destination.exists()

    output_book = load_workbook(destination, read_only=True, data_only=True)
    try:
        rows = list(output_book["INVENTARIO"].iter_rows(values_only=True))
        records = rows[1:]
        assert [row[1] for row in records] == ordered_ips
        assert [row[3] for row in records] == [family for family, _ in _FAMILIES]
        assert [row[5] for row in records] == [
            f"Synthetic {family} model" for family, _ in _FAMILIES
        ]
        exported_content = "\n".join(
            str(value) for row in rows for value in row if value is not None
        )
        assert all(password not in exported_content for password in secrets)
        assert "synthetic-user" not in exported_content
    finally:
        output_book.close()

    for ip in ordered_ips:
        target_calls = [item for item in calls if item[0] == ip]
        assert any(item[1:] == ("GetCapabilities", False) for item in target_calls)
        assert any(item[1:] == ("GetDeviceInformation", True) for item in target_calls)
