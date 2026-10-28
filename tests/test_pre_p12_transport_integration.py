"""Offline workflow integration checks for the real Zeep/Requests transport."""
from __future__ import annotations

import ssl
import threading
import zipfile
from collections.abc import Callable
from contextlib import ExitStack
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from io import BytesIO
from pathlib import Path
from typing import Self
from urllib.parse import urlsplit

import pytest
from openpyxl import Workbook, load_workbook
from PIL import Image
from requests.adapters import HTTPAdapter

from cam_scanner.application.collection_context import CollectionContext
from cam_scanner.application.inventory_service import InventoryService
from cam_scanner.application.multi_workflow import MultiWorkflow
from cam_scanner.application.single_workflow import SingleWorkflow
from cam_scanner.cameras.manufacturers.hikvision import HikvisionAdapter
from cam_scanner.cameras.onvif.adapter import GenericOnvifCollector
from cam_scanner.cameras.preauth import OnvifPreAuthDiscoverer, PreAuthEvidence
from cam_scanner.cameras.registry import AdapterRegistry
from cam_scanner.cameras.strategy import VendorFirstCollector
from cam_scanner.cameras.transport import ControlledOnvifClient, TargetTransportPolicy
from cam_scanner.config.models import EffectiveConfig, TransportProfile, TransportRule
from cam_scanner.config.transport_plan import select_transport
from cam_scanner.domain.enums import CollectionMethod, CollectionStatus, MainMenuChoice
from cam_scanner.domain.models import CameraResult, CameraTarget, CollectionReport
from cam_scanner.excel.reader import ExcelInputRow
from cam_scanner.snapshot.service import (
    SnapshotFailure,
    SnapshotMethod,
    SnapshotResult,
    SnapshotService,
)

SOAP_DEVICE_INFORMATION = (
    b'<s:Envelope xmlns:s="http://www.w3.org/2003/05/soap-envelope" '
    b'xmlns:tds="http://www.onvif.org/ver10/device/wsdl"><s:Body>'
    b'<tds:GetDeviceInformationResponse><tds:Manufacturer>Axis</tds:Manufacturer>'
    b'<tds:Model>M1</tds:Model><tds:FirmwareVersion>1.0</tds:FirmwareVersion>'
    b'<tds:SerialNumber>S1</tds:SerialNumber><tds:HardwareId>H1</tds:HardwareId>'
    b'</tds:GetDeviceInformationResponse></s:Body></s:Envelope>'
)


class _Server:
    def __init__(self, handler: type[BaseHTTPRequestHandler]) -> None:
        self.httpd = ThreadingHTTPServer(("127.0.0.1", 0), handler)
        self.thread = threading.Thread(target=self.httpd.serve_forever, daemon=True)

    @property
    def port(self) -> int:
        return self.httpd.server_port

    def __enter__(self) -> Self:
        self.thread.start()
        return self

    def __exit__(self, *_exc: object) -> None:
        self.httpd.shutdown()
        self.httpd.server_close()
        self.thread.join(timeout=2)


class _HttpsServer(_Server):
    def __init__(self, handler: type[BaseHTTPRequestHandler]) -> None:
        super().__init__(handler)
        tls = ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
        tls.load_cert_chain(
            Path(__file__).with_name("tls_test_server_cardb.pem"),
            Path(__file__).with_name("tls_test_server_cardb.key"),
        )
        self.httpd.socket = tls.wrap_socket(self.httpd.socket, server_side=True)


def _soap_snapshot_handler(addresses: dict[str, str], observed: list[tuple[str, bytes]]):
    class Handler(BaseHTTPRequestHandler):
        def do_POST(self) -> None:
            body = self.rfile.read(int(self.headers.get("Content-Length", "0")))
            observed.append((self.path, body))
            if b"GetCapabilities" in body:
                operation = (
                    "<tds:GetCapabilitiesResponse><tds:Capabilities>"
                    "<tt:Media><tt:XAddr>{media}</tt:XAddr></tt:Media>"
                    "</tds:Capabilities></tds:GetCapabilitiesResponse>"
                ).format(media=addresses["media"])
            elif b"GetDeviceInformation" in body:
                operation = (
                    "<tds:GetDeviceInformationResponse>"
                    "<tds:Manufacturer>Axis</tds:Manufacturer>"
                    "<tds:Model>M1</tds:Model><tds:FirmwareVersion>1.0</tds:FirmwareVersion>"
                    "<tds:SerialNumber>S1</tds:SerialNumber><tds:HardwareId>H1</tds:HardwareId>"
                    "</tds:GetDeviceInformationResponse>"
                )
            elif b"GetProfiles" in body:
                operation = (
                    '<trt:GetProfilesResponse><trt:Profiles token="profile-1" fixed="true">'
                    "<tt:Name>profile</tt:Name></trt:Profiles></trt:GetProfilesResponse>"
                )
            elif b"GetSnapshotUri" in body:
                operation = (
                    "<trt:GetSnapshotUriResponse><trt:MediaUri>"
                    "<tt:Uri>{snapshot}</tt:Uri><tt:InvalidAfterConnect>false</tt:InvalidAfterConnect>"
                    "<tt:InvalidAfterReboot>false</tt:InvalidAfterReboot><tt:Timeout>PT1M</tt:Timeout>"
                    "</trt:MediaUri></trt:GetSnapshotUriResponse>"
                ).format(snapshot=addresses["snapshot"])
            else:
                self.send_response(500)
                self.send_header("Content-Length", "0")
                self.end_headers()
                return
            envelope = (
                '<s:Envelope xmlns:s="http://www.w3.org/2003/05/soap-envelope" '
                'xmlns:tds="http://www.onvif.org/ver10/device/wsdl" '
                'xmlns:trt="http://www.onvif.org/ver10/media/wsdl" '
                'xmlns:tt="http://www.onvif.org/ver10/schema"><s:Body>'
                f"{operation}</s:Body></s:Envelope>"
            ).encode()
            self.send_response(200)
            self.send_header("Content-Type", "application/soap+xml; charset=utf-8")
            self.send_header("Content-Length", str(len(envelope)))
            self.end_headers()
            self.wfile.write(envelope)

        def do_GET(self) -> None:
            observed.append((self.path, b"GET"))
            if self.path != "/snapshot":
                self.send_response(404)
                self.send_header("Content-Length", "0")
                self.end_headers()
                return
            image = b"\xff\xd8\xff\xd9"
            self.send_response(200)
            self.send_header("Content-Type", "image/jpeg")
            self.send_header("Content-Length", str(len(image)))
            self.end_headers()
            self.wfile.write(image)

        def log_message(self, *_args: object) -> None:
            pass

    return Handler


def _trust_test_ca_for_ports(monkeypatch, ports: set[int]) -> None:
    """Trust the fixture CA only for selected loopback requests; keep real TLS checks."""
    original_send = HTTPAdapter.send
    ca_bundle = str(Path(__file__).with_name("tls_test_ca_cardb.pem"))

    def send(adapter, request, **kwargs):
        parsed = urlsplit(request.url)
        if parsed.scheme == "https" and parsed.hostname == "127.0.0.1" and parsed.port in ports:
            assert kwargs.get("verify") is True
            kwargs["verify"] = ca_bundle
        return original_send(adapter, request, **kwargs)

    monkeypatch.setattr(HTTPAdapter, "send", send)


def _https_rules(port: int, *, components: tuple[str, ...]) -> tuple[TransportRule, ...]:
    return tuple(
        TransportRule("127.0.0.1", component, "https", None, port)
        for component in components
    )


def _soap_handler(
    *,
    on_post: Callable[[BaseHTTPRequestHandler], None] | None = None,
) -> type[BaseHTTPRequestHandler]:
    class Handler(BaseHTTPRequestHandler):
        def do_POST(self) -> None:
            self.rfile.read(int(self.headers.get("Content-Length", "0")))
            if on_post is not None:
                on_post(self)
                return
            self.send_response(200)
            self.send_header("Content-Type", "application/soap+xml")
            self.send_header("Content-Length", str(len(SOAP_DEVICE_INFORMATION)))
            self.end_headers()
            self.wfile.write(SOAP_DEVICE_INFORMATION)

        def log_message(self, *_args: object) -> None:
            pass

    return Handler


def _negative_workflow_handler(
    scenario: str,
    addresses: dict[str, str],
    observed: list[tuple[str, bytes]],
    redirect_status: int = 302,
) -> type[BaseHTTPRequestHandler]:
    class Handler(BaseHTTPRequestHandler):
        def do_POST(self) -> None:
            body = self.rfile.read(int(self.headers.get("Content-Length", "0")))
            observed.append((self.path, body))
            if scenario == "preauth_redirect":
                self.send_response(redirect_status)
                self.send_header("Location", addresses["redirect"])
                self.send_header("Content-Length", "0")
                self.end_headers()
                return
            if b"GetCapabilities" in body:
                operation = (
                    "<tds:GetCapabilitiesResponse><tds:Capabilities>"
                    "<tt:Media><tt:XAddr>{media}</tt:XAddr></tt:Media>"
                    "</tds:Capabilities></tds:GetCapabilitiesResponse>"
                ).format(media=addresses["media"])
            elif b"GetDeviceInformation" in body:
                operation = (
                    "<tds:GetDeviceInformationResponse>"
                    "<tds:Manufacturer>Hikvision</tds:Manufacturer>"
                    "<tds:Model>DS-KD8003</tds:Model><tds:FirmwareVersion>1.0</tds:FirmwareVersion>"
                    "<tds:SerialNumber>S1</tds:SerialNumber><tds:HardwareId>H1</tds:HardwareId>"
                    "</tds:GetDeviceInformationResponse>"
                )
            elif b"GetProfiles" in body:
                operation = (
                    '<trt:GetProfilesResponse><trt:Profiles token="profile-1" fixed="true">'
                    "<tt:Name>profile</tt:Name></trt:Profiles></trt:GetProfilesResponse>"
                )
            elif b"GetSnapshotUri" in body:
                operation = (
                    "<trt:GetSnapshotUriResponse><trt:MediaUri>"
                    "<tt:Uri>{snapshot}</tt:Uri><tt:InvalidAfterConnect>false</tt:InvalidAfterConnect>"
                    "<tt:InvalidAfterReboot>false</tt:InvalidAfterReboot><tt:Timeout>PT1M</tt:Timeout>"
                    "</trt:MediaUri></trt:GetSnapshotUriResponse>"
                ).format(snapshot=addresses["snapshot"])
            else:
                self.send_response(500)
                self.send_header("Content-Length", "0")
                self.end_headers()
                return
            envelope = (
                '<s:Envelope xmlns:s="http://www.w3.org/2003/05/soap-envelope" '
                'xmlns:tds="http://www.onvif.org/ver10/device/wsdl" '
                'xmlns:trt="http://www.onvif.org/ver10/media/wsdl" '
                'xmlns:tt="http://www.onvif.org/ver10/schema"><s:Body>'
                f"{operation}</s:Body></s:Envelope>"
            ).encode()
            self.send_response(200)
            self.send_header("Content-Type", "application/soap+xml; charset=utf-8")
            self.send_header("Content-Length", str(len(envelope)))
            self.end_headers()
            self.wfile.write(envelope)

        def do_GET(self) -> None:
            observed.append((self.path, b"GET"))
            if self.path == "/ISAPI/System/deviceInfo":
                body = (
                    b"<DeviceInfo><manufacturer>Hikvision</manufacturer>"
                    b"<model>DS-KD8003</model></DeviceInfo>"
                )
                self.send_response(200)
                self.send_header("Content-Type", "application/xml")
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)
                return
            if self.path == "/snapshot" and scenario == "snapshot_redirect":
                self.send_response(redirect_status)
                self.send_header("Location", addresses["redirect"])
                self.send_header("Content-Length", "0")
                self.end_headers()
                return
            if self.path != "/snapshot":
                self.send_response(404)
                self.send_header("Content-Length", "0")
                self.end_headers()
                return
            image = b"\xff\xd8\xff\xd9"
            self.send_response(200)
            self.send_header("Content-Type", "image/jpeg")
            self.send_header("Content-Length", str(len(image)))
            self.end_headers()
            self.wfile.write(image)

        def log_message(self, *_args: object) -> None:
            pass

    return Handler


def _sink_handler(observed: list[str]) -> type[BaseHTTPRequestHandler]:
    class Handler(BaseHTTPRequestHandler):
        def do_GET(self) -> None:
            observed.append(self.path)
            self.send_response(200)
            self.send_header("Content-Length", "0")
            self.end_headers()

        def do_POST(self) -> None:
            observed.append(self.path)
            self.send_response(200)
            self.send_header("Content-Length", "0")
            self.end_headers()

        def log_message(self, *_args: object) -> None:
            pass

    return Handler


def _real_hikvision_workflow_inventory(config: EffectiveConfig) -> InventoryService:
    registry = AdapterRegistry()
    registry.register("Hikvision", HikvisionAdapter())
    collector = VendorFirstCollector(
        onvif_collector=GenericOnvifCollector(),
        registry=registry,
        discoverer=OnvifPreAuthDiscoverer(),
    )
    return InventoryService(collector, config=config)


class _RealZeepCollector:
    """Adapt a real Device service operation to InventoryService's report contract."""

    def collect_report(self, target: CameraTarget, *, context: object) -> CollectionReport:
        assert context.auth.begin("ONVIF", "primary")
        response = ControlledOnvifClient(context, authenticated=True).devicemgmt().GetDeviceInformation()
        assert response.Manufacturer == "Axis"
        result = CameraResult(
            ip=target.ip,
            status=CollectionStatus.SUCCESS,
            manufacturer=response.Manufacturer,
            model=response.Model,
            serial=response.SerialNumber,
            firmware=response.FirmwareVersion,
            hardware_id=response.HardwareId,
            collection_method=CollectionMethod.ONVIF,
        )
        return CollectionReport.from_camera_result(result)


class _NoSnapshot:
    def capture_one(self, *_args: object, **_kwargs: object) -> SnapshotResult:
        return SnapshotResult(failure=SnapshotFailure.UNSUPPORTED)


class _EmptyDiscovery:
    def discover_with_report(
        self, _target: CameraTarget, *, context: object
    ) -> tuple[PreAuthEvidence, CollectionReport]:
        del context
        return PreAuthEvidence(), CollectionReport()


def _real_onvif_inventory(port: int) -> InventoryService:
    config = EffectiveConfig(transport_rules=_https_rules(
        port, components=("onvif_device",)
    ))
    collector = VendorFirstCollector(
        onvif_collector=GenericOnvifCollector(),
        discoverer=_EmptyDiscovery(),
    )
    return InventoryService(collector, config=config)


def _auth_rejected_handler(requests: list[tuple[str, str]]) -> type[BaseHTTPRequestHandler]:
    class Handler(BaseHTTPRequestHandler):
        def do_POST(self) -> None:
            self.rfile.read(int(self.headers.get("Content-Length", "0")))
            requests.append(("POST", self.path))
            self.send_response(401)
            self.send_header("Content-Length", "0")
            self.end_headers()

        def do_GET(self) -> None:
            requests.append(("GET", self.path))
            self.send_response(401)
            self.send_header("Content-Length", "0")
            self.end_headers()

        def log_message(self, *_args: object) -> None:
            pass

    return Handler


class _SingleWorkflowUI:
    def __init__(self, target: CameraTarget) -> None:
        self._target = target
        self.report: CollectionReport | None = None
        self.snapshot: SnapshotResult | None = None

    def prompt_camera_target(self) -> CameraTarget:
        return self._target

    def render_report(
        self,
        report: CollectionReport,
        *,
        ip: str,
        duration: float,
        snapshot: SnapshotResult | None,
    ) -> None:
        del ip, duration
        self.report = report
        self.snapshot = snapshot

    def render_result(self, *_args: object) -> None:
        raise AssertionError("the authenticated inventory path should return a report")

    def post_query_menu(self) -> MainMenuChoice:
        return MainMenuChoice.EXIT


class _MultiWorkflowUI:
    def __init__(self, source: Path, destination: Path) -> None:
        self.source = source
        self.destination = destination
        self.completed: tuple[Path, int, int] | None = None
        self.errors: list[str] = []

    def prompt_multi_input_path(self) -> str:
        return str(self.source)

    def prompt_multi_output_path(self, _default: str) -> str:
        return str(self.destination)

    def show_multi_complete(self, path: Path, total: int, *, failed: int) -> None:
        self.completed = (path, total, failed)

    def show_multi_error(self, message: str) -> None:
        self.errors.append(message)


class _QuietProgress:
    def start(self, _total: int) -> None:
        pass

    def advance(self, _ip: str, _status: str) -> None:
        pass

    def finish(self) -> None:
        pass


def _service(port: int) -> InventoryService:
    config = EffectiveConfig(transport_rules=_https_rules(
        port, components=("onvif_device",)
    ))
    return InventoryService(_RealZeepCollector(), config=config)


def _real_https_workflow_inventory(port: int) -> InventoryService:
    config = EffectiveConfig(
        transport_rules=_https_rules(
            port,
            components=("onvif_device", "onvif_media", "snapshot_onvif"),
        )
    )
    return InventoryService(
        VendorFirstCollector(discoverer=OnvifPreAuthDiscoverer()),
        config=config,
    )


def _jpeg_fixture() -> bytes:
    output = BytesIO()
    Image.new("RGB", (16, 8), color=(30, 60, 90)).save(output, format="JPEG")
    return output.getvalue()


def test_single_inventory_path_uses_real_zeep_and_exact_target_port(monkeypatch) -> None:
    observed: list[tuple[str, bytes]] = []

    def record(handler: BaseHTTPRequestHandler) -> None:
        observed.append((handler.path, handler.headers.get("Content-Type", "").encode()))
        handler.send_response(200)
        handler.send_header("Content-Type", "application/soap+xml")
        handler.send_header("Content-Length", str(len(SOAP_DEVICE_INFORMATION)))
        handler.end_headers()
        handler.wfile.write(SOAP_DEVICE_INFORMATION)

    with _HttpsServer(_soap_handler(on_post=record)) as server:
        _trust_test_ca_for_ports(monkeypatch, {server.port})
        service = _service(server.port)
        target = CameraTarget("127.0.0.1", "offline-user", "offline-password")
        report, context = service.collect_execution(target)
        try:
            result = report.to_camera_result(target.ip)
            assert result.status is CollectionStatus.SUCCESS
            assert (result.manufacturer, result.model, result.serial) == ("Axis", "M1", "S1")
            assert len(observed) == 1
            assert observed[0][0] == "/onvif/device_service"
            assert observed[0][1].startswith(b"application/soap+xml")
            assert context.auth.physical_exchanges == 1
            assert len(result.__dataclass_fields__) == 14
        finally:
            context.close()


def test_multi_row_real_zeep_redirect_does_not_reach_unapproved_loopback_port(monkeypatch) -> None:
    redirected_requests: list[str] = []

    def record_redirect(handler: BaseHTTPRequestHandler) -> None:
        handler.send_response(302)
        handler.send_header("Location", f"http://127.0.0.1:{destination.port}/other")
        handler.send_header("Content-Length", "0")
        handler.end_headers()

    class DestinationHandler(BaseHTTPRequestHandler):
        def do_POST(self) -> None:
            redirected_requests.append(self.path)
            self.send_response(200)
            self.send_header("Content-Length", "0")
            self.end_headers()

        def log_message(self, *_args: object) -> None:
            pass

    with _Server(DestinationHandler) as destination, _HttpsServer(
        _soap_handler(on_post=record_redirect)
    ) as source:
        _trust_test_ca_for_ports(monkeypatch, {source.port})
        service = _service(source.port)
        row = ExcelInputRow(
            row_id=2,
            ip="127.0.0.1",
            username="offline-user",
            password="offline-password",
        )
        workflow = MultiWorkflow(
            ui=object(),  # _process_row only uses the UI on batch-level failures.
            inventory_service=service,
            snapshot_service=_NoSnapshot(),  # type: ignore[arg-type]
            max_workers=1,
        )

        output = workflow._process_row(row)

    assert redirected_requests == []
    assert output.result.ip == "127.0.0.1"
    assert output.result.status is CollectionStatus.FAILED


def test_single_workflow_stops_all_authenticated_io_after_real_onvif_401(monkeypatch) -> None:
    requests: list[tuple[str, str]] = []
    target = CameraTarget("127.0.0.1", "offline-user", "offline-password")

    with _HttpsServer(_auth_rejected_handler(requests)) as server:
        _trust_test_ca_for_ports(monkeypatch, {server.port})
        ui = _SingleWorkflowUI(target)
        SingleWorkflow(
            ui=ui,  # type: ignore[arg-type]
            inventory_service=_real_onvif_inventory(server.port),
            snapshot_service=SnapshotService(),
        ).run()

    assert ui.report is not None
    assert ui.snapshot is not None
    assert ui.snapshot.failure is SnapshotFailure.AUTH_REJECTED
    assert requests == [("POST", "/onvif/device_service")]


def test_multi_row_stops_all_authenticated_io_after_real_onvif_401(monkeypatch) -> None:
    requests: list[tuple[str, str]] = []

    with _HttpsServer(_auth_rejected_handler(requests)) as server:
        _trust_test_ca_for_ports(monkeypatch, {server.port})
        workflow = MultiWorkflow(
            ui=object(),  # _process_row only uses the UI on batch-level failures.
            inventory_service=_real_onvif_inventory(server.port),
            snapshot_service=SnapshotService(),
            max_workers=1,
        )
        result = workflow._process_row(
            ExcelInputRow(
                row_id=2,
                ip="127.0.0.1",
                username="offline-user",
                password="offline-password",
            )
        )

    assert result.result.status is not CollectionStatus.SUCCESS
    assert requests == [("POST", "/onvif/device_service")]


def test_trusted_https_uses_real_zeep_and_snapshot_transport(monkeypatch) -> None:
    observed: list[tuple[str, bytes]] = []
    addresses = {"media": "", "snapshot": ""}

    with _HttpsServer(_soap_snapshot_handler(addresses, observed)) as server:
        addresses["media"] = f"https://127.0.0.1:{server.port}/onvif/media_service"
        addresses["snapshot"] = f"https://127.0.0.1:{server.port}/snapshot"
        _trust_test_ca_for_ports(monkeypatch, {server.port})
        target = CameraTarget("127.0.0.1", "offline-user", "offline-password")
        config = EffectiveConfig(transport_rules=_https_rules(
            server.port,
            components=("onvif_device", "onvif_media", "snapshot_onvif"),
        ))
        context = CollectionContext(
            target, config, transport_factory=TargetTransportPolicy
        )
        assert context.auth.begin("ONVIF", "primary")
        try:
            device = ControlledOnvifClient(context, authenticated=True).devicemgmt()
            capabilities = device.GetCapabilities(Category="All")
            assert capabilities.Media.XAddr == addresses["media"]
            details = device.GetDeviceInformation()
            assert details.Manufacturer == "Axis"
            assert details.SerialNumber == "S1"

            result = SnapshotService(
                image_processor=lambda _raw: (b"thumbnail", 320, 180)
            ).capture_one(target, manufacturer="Axis", model="M1", context=context)

            assert result.successful
            assert result.method is SnapshotMethod.ONVIF_HTTP
            assert result.attempts[-1].success
            assert any(path == "/snapshot" and body == b"GET" for path, body in observed)
            assert sum(path == "/onvif/device_service" for path, _ in observed) == 2
            assert sum(path == "/onvif/media_service" for path, _ in observed) == 2
            assert context.auth.physical_exchanges == 5
        finally:
            assert context.close()


def test_single_workflow_runs_real_https_preauth_inventory_and_snapshot(monkeypatch) -> None:
    observed: list[tuple[str, bytes]] = []
    addresses = {"media": "", "snapshot": ""}

    with _HttpsServer(_soap_snapshot_handler(addresses, observed)) as server:
        addresses["media"] = f"https://127.0.0.1:{server.port}/onvif/media_service"
        addresses["snapshot"] = f"https://127.0.0.1:{server.port}/snapshot"
        _trust_test_ca_for_ports(monkeypatch, {server.port})
        target = CameraTarget("127.0.0.1", "offline-user", "offline-password")
        ui = _SingleWorkflowUI(target)
        SingleWorkflow(
            ui,
            _real_https_workflow_inventory(server.port),
            SnapshotService(image_processor=lambda _raw: (_jpeg_fixture(), 16, 8)),
        ).run()

    assert ui.report is not None
    result = ui.report.to_camera_result(target.ip)
    assert result.status is CollectionStatus.SUCCESS
    assert (result.manufacturer, result.model, result.serial) == ("Axis", "M1", "S1")
    assert ui.snapshot is not None and ui.snapshot.successful
    soap_bodies = [body for path, body in observed if path != "/snapshot"]
    assert any(b"GetCapabilities" in body for body in soap_bodies)
    assert any(b"GetDeviceInformation" in body for body in soap_bodies)
    assert any(b"GetProfiles" in body for body in soap_bodies)
    assert any(b"GetSnapshotUri" in body for body in soap_bodies)
    assert b"offline-user" not in soap_bodies[0]
    assert b"offline-password" not in soap_bodies[0]
    assert any(path == "/snapshot" and body == b"GET" for path, body in observed)


def test_multi_workflow_runs_real_https_preauth_inventory_and_snapshot(
    tmp_path: Path, monkeypatch
) -> None:
    observed: list[tuple[str, bytes]] = []
    addresses = {"media": "", "snapshot": ""}
    source = tmp_path / "cameras.xlsx"
    destination = tmp_path / "inventory.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(("IP", "USERNAME", "PASSWORD"))
    worksheet.append(("127.0.0.1", "offline-user", "offline-password"))
    workbook.save(source)
    workbook.close()

    with _HttpsServer(_soap_snapshot_handler(addresses, observed)) as server:
        addresses["media"] = f"https://127.0.0.1:{server.port}/onvif/media_service"
        addresses["snapshot"] = f"https://127.0.0.1:{server.port}/snapshot"
        _trust_test_ca_for_ports(monkeypatch, {server.port})
        ui = _MultiWorkflowUI(source, destination)
        MultiWorkflow(
            ui,
            _real_https_workflow_inventory(server.port),
            SnapshotService(image_processor=lambda _raw: (_jpeg_fixture(), 16, 8)),
            progress_reporter=_QuietProgress(),
            max_workers=1,
        ).run()

    assert ui.errors == []
    assert ui.completed == (destination, 1, 0)
    assert destination.is_file()
    output = load_workbook(destination, data_only=True)
    inventory = output["INVENTARIO"]
    assert inventory.cell(2, 11).value == CollectionStatus.SUCCESS.value
    assert inventory.cell(2, 4).value == "Axis"
    assert inventory.cell(2, 16).value == "Incluído"
    output.close()
    soap_bodies = [body for path, body in observed if path != "/snapshot"]
    assert any(b"GetCapabilities" in body for body in soap_bodies)
    assert any(b"GetDeviceInformation" in body for body in soap_bodies)
    assert b"offline-user" not in soap_bodies[0]
    assert b"offline-password" not in soap_bodies[0]
    assert any(path == "/snapshot" and body == b"GET" for path, body in observed)
    with zipfile.ZipFile(destination) as archive:
        workbook_bytes = b"".join(archive.read(name) for name in archive.namelist())
    assert b"offline-user" not in workbook_bytes
    assert b"offline-password" not in workbook_bytes


@pytest.mark.parametrize("workflow_kind", ("single", "multi"))
@pytest.mark.parametrize(
    "scenario,expected_failure,redirect_status",
    (
        ("preauth_tls", SnapshotFailure.TLS_ERROR, None),
        *tuple(
            ("preauth_redirect", SnapshotFailure.AUTH_REJECTED, status)
            for status in (301, 302, 303, 307, 308)
        ),
        ("media_unapproved", SnapshotFailure.TARGET_REJECTED, None),
        ("snapshot_unapproved", SnapshotFailure.TARGET_REJECTED, None),
        *tuple(
            ("snapshot_redirect", SnapshotFailure.TARGET_REJECTED, status)
            for status in (301, 302, 303, 307, 308)
        ),
        ("snapshot_tls", SnapshotFailure.TLS_ERROR, None),
    ),
)
def test_full_workflows_fail_closed_for_tls_redirects_and_unapproved_endpoints(
    workflow_kind: str,
    scenario: str,
    expected_failure: SnapshotFailure,
    redirect_status: int | None,
    tmp_path: Path,
    monkeypatch,
) -> None:
    """Exercise transport failures through the actual SINGLE and MULTI entry points."""
    observed: list[tuple[str, bytes]] = []
    sink_requests: list[str] = []
    vendor_snapshot_requests: list[str] = []
    rtsp_calls: list[str] = []
    addresses = {"media": "", "snapshot": "", "redirect": ""}
    target = CameraTarget("127.0.0.1", "offline-user", "offline-password")

    class RtspSpy:
        def capture(self, *_args: object, **_kwargs: object) -> bytes:
            rtsp_calls.append("capture")
            return _jpeg_fixture()

    with ExitStack() as stack:
        sink = stack.enter_context(_Server(_sink_handler(sink_requests)))
        vendor_snapshot = stack.enter_context(
            _Server(_sink_handler(vendor_snapshot_requests))
        )
        addresses["redirect"] = f"http://127.0.0.1:{sink.port}/redirected"

        bad_snapshot = None
        control = stack.enter_context(
            _HttpsServer(_negative_workflow_handler(
                scenario, addresses, observed, redirect_status or 302
            ))
        )
        if scenario == "snapshot_tls":
            bad_snapshot = stack.enter_context(
                _HttpsServer(_negative_workflow_handler(
                    scenario, addresses, observed, redirect_status or 302
                ))
            )

        addresses["media"] = f"https://127.0.0.1:{control.port}/onvif/media_service"
        addresses["snapshot"] = f"https://127.0.0.1:{control.port}/snapshot"
        if scenario == "media_unapproved":
            addresses["media"] = f"http://127.0.0.1:{sink.port}/onvif/media_service"
        elif scenario == "snapshot_unapproved":
            addresses["snapshot"] = f"https://127.0.0.1:{sink.port}/snapshot"
        elif scenario == "snapshot_tls":
            assert bad_snapshot is not None
            addresses["snapshot"] = f"https://127.0.0.1:{bad_snapshot.port}/snapshot"

        if scenario != "preauth_tls":
            _trust_test_ca_for_ports(monkeypatch, {control.port})

        components = ("onvif_device", "onvif_media", "vendor_hikvision")
        rules = _https_rules(control.port, components=components)
        snapshot_port = (
            bad_snapshot.port
            if scenario == "snapshot_tls" and bad_snapshot is not None
            else control.port
        )
        rules += _https_rules(snapshot_port, components=("snapshot_onvif",))
        rules += (
            TransportRule(
                "127.0.0.1", "snapshot_vendor", "http", None, vendor_snapshot.port
            ),
        )
        config = EffectiveConfig(transport_rules=rules)
        inventory = _real_hikvision_workflow_inventory(config)
        snapshots = SnapshotService(
            rtsp_provider=RtspSpy(),
            image_processor=lambda _raw: (_jpeg_fixture(), 16, 8),
        )

        if workflow_kind == "single":
            ui = _SingleWorkflowUI(target)
            SingleWorkflow(ui, inventory, snapshots).run()
            assert ui.report is not None
            assert ui.snapshot is not None
            snapshot = ui.snapshot
            assert any(
                attempt.error_category == "TLS_ERROR"
                for attempt in ui.report.attempts
            ) is (scenario == "preauth_tls")
        else:
            source = tmp_path / "negative-cameras.xlsx"
            destination = tmp_path / "negative-inventory.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(("IP", "USERNAME", "PASSWORD"))
            worksheet.append((target.ip, target.username, target.password))
            workbook.save(source)
            workbook.close()
            ui = _MultiWorkflowUI(source, destination)
            workflow = MultiWorkflow(
                ui,
                inventory,
                snapshots,
                progress_reporter=_QuietProgress(),
                max_workers=1,
            )
            workflow.run()
            assert ui.errors == []
            assert ui.completed == (
                destination,
                1,
                int(workflow.completed_rows[0].result.status is CollectionStatus.FAILED),
            )
            assert destination.is_file()
            assert len(workflow.completed_rows) == 1
            snapshot = workflow.completed_rows[0].snapshot

    assert snapshot is not None
    assert snapshot.failure is expected_failure, (
        f"{workflow_kind=} {scenario=} actual={snapshot.failure!r}"
    )
    assert sink_requests == []
    assert vendor_snapshot_requests == []
    assert rtsp_calls == []
    if scenario == "preauth_tls":
        assert observed == []
    else:
        assert any(path == "/onvif/device_service" for path, _ in observed)


@pytest.mark.parametrize("scheme", ("http", "https"))
def test_real_target_transport_ignores_hostile_proxies_netrc_and_ambient_credentials(
    scheme: str,
    tmp_path: Path,
    monkeypatch,
) -> None:
    """A real product request stays direct and anonymous despite hostile process env."""
    target_requests: list[dict[str, str]] = []
    proxy_requests: list[dict[str, str]] = []

    class TargetHandler(BaseHTTPRequestHandler):
        def do_GET(self) -> None:
            target_requests.append(dict(self.headers.items()))
            body = b"offline transport response"
            self.send_response(200)
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def log_message(self, *_args: object) -> None:
            pass

    class HostileProxyHandler(BaseHTTPRequestHandler):
        def _record(self) -> None:
            proxy_requests.append({
                "method": self.command,
                "path": self.path,
                **dict(self.headers.items()),
            })
            self.send_response(502)
            self.send_header("Content-Length", "0")
            self.end_headers()

        do_GET = _record
        do_POST = _record
        do_CONNECT = _record

        def log_message(self, *_args: object) -> None:
            pass

    with ExitStack() as stack:
        proxy = stack.enter_context(_Server(HostileProxyHandler))
        server_type = _HttpsServer if scheme == "https" else _Server
        target_server = stack.enter_context(server_type(TargetHandler))
        if scheme == "https":
            _trust_test_ca_for_ports(monkeypatch, {target_server.port})

        netrc_user = "ambient-netrc-user"
        netrc_password = "ambient-netrc-password"
        proxy_user = "ambient-proxy-user"
        proxy_password = "ambient-proxy-password"
        netrc_path = tmp_path / "hostile.netrc"
        netrc_path.write_text(
            f"machine 127.0.0.1 login {netrc_user} password {netrc_password}\n",
            encoding="utf-8",
        )
        proxy_url = (
            f"http://{proxy_user}:{proxy_password}@127.0.0.1:{proxy.port}"
        )
        for name in (
            "HTTP_PROXY", "HTTPS_PROXY", "ALL_PROXY",
            "http_proxy", "https_proxy", "all_proxy",
        ):
            monkeypatch.setenv(name, proxy_url)
        for name in ("NO_PROXY", "no_proxy"):
            monkeypatch.setenv(name, "")
        monkeypatch.setenv("NETRC", str(netrc_path))
        monkeypatch.setenv("USER", netrc_user)
        monkeypatch.setenv("USERNAME", netrc_user)
        monkeypatch.setenv("LOGNAME", netrc_user)
        monkeypatch.setenv("PASSWORD", netrc_password)
        monkeypatch.setenv("CAMERA_USERNAME", netrc_user)
        monkeypatch.setenv("CAMERA_PASSWORD", netrc_password)

        target = CameraTarget("127.0.0.1", "configured-user", "configured-password")
        if scheme == "http":
            profile = TransportProfile(
                id="OFFLINE_TEST_HTTP",
                targets=(target.ip,),
                component="vendor_axis",
                protocol="http",
                port=target_server.port,
                auth_mode="PUBLIC_NONE",
                authorization_ref="OFFLINE_TEST_AUTH",
                http_risk_accepted=True,
            )
            config = select_transport(
                EffectiveConfig(transport_profiles=(profile,)),
                (target,), {target.ip: "vendor_axis"}, profile.id,
            )
        else:
            config = EffectiveConfig(transport_rules=(TransportRule(
                target.ip, "vendor_axis", scheme, None, target_server.port
            ),))
        context = CollectionContext(
            target, config, transport_factory=TargetTransportPolicy
        )
        try:
            url = f"{scheme}://127.0.0.1:{target_server.port}/anonymous-probe"
            response = context.transport.request(
                "vendor_axis", "GET", url, source="NATIVE", authenticated=False
            )
            assert response.status_code == 200
            response.close()

            assert len(target_requests) == 1
            assert not any(
                key.casefold() in {"authorization", "proxy-authorization"}
                for key in target_requests[0]
            )
            assert proxy_requests == []
            captured = repr((target_requests, proxy_requests)).casefold()
            for secret in (
                netrc_user, netrc_password, proxy_user, proxy_password,
                "configured-user", "configured-password",
            ):
                assert secret.casefold() not in captured
            session = context.transport._sessions[("vendor_axis", "NATIVE", False)]
            assert session.trust_env is False
            assert session.proxies == {}
        finally:
            context.close()


def test_untrusted_snapshot_tls_does_not_downgrade_to_explicit_vendor_http(monkeypatch) -> None:
    control_requests: list[tuple[str, bytes]] = []
    bad_snapshot_requests: list[str] = []
    vendor_requests: list[str] = []
    addresses = {"media": "", "snapshot": ""}

    class BadSnapshotHandler(BaseHTTPRequestHandler):
        def do_GET(self) -> None:
            bad_snapshot_requests.append(self.path)
            self.send_response(200)
            self.send_header("Content-Length", "4")
            self.end_headers()
            self.wfile.write(b"jpeg")

        def log_message(self, *_args: object) -> None:
            pass

    class VendorHandler(BaseHTTPRequestHandler):
        def do_GET(self) -> None:
            vendor_requests.append(self.path)
            self.send_response(200)
            self.send_header("Content-Length", "4")
            self.end_headers()
            self.wfile.write(b"jpeg")

        def log_message(self, *_args: object) -> None:
            pass

    class RtspSpy:
        def capture(self, *_args: object, **_kwargs: object) -> bytes:
            raise AssertionError("RTSP must not run after TLS validation failure")

    with _HttpsServer(BadSnapshotHandler) as bad_snapshot, _Server(
        VendorHandler
    ) as vendor, _HttpsServer(_soap_snapshot_handler(addresses, control_requests)) as control:
        addresses["media"] = f"https://127.0.0.1:{control.port}/onvif/media_service"
        addresses["snapshot"] = f"https://127.0.0.1:{bad_snapshot.port}/snapshot"
        _trust_test_ca_for_ports(monkeypatch, {control.port})
        target = CameraTarget("127.0.0.1", "offline-user", "offline-password")
        rules = _https_rules(
            control.port, components=("onvif_device", "onvif_media")
        ) + _https_rules(bad_snapshot.port, components=("snapshot_onvif",)) + (
            TransportRule("127.0.0.1", "snapshot_vendor", "http", None, vendor.port),
        )
        context = CollectionContext(
            target,
            EffectiveConfig(transport_rules=rules),
            transport_factory=TargetTransportPolicy,
        )
        assert context.auth.begin("VENDOR_HIKVISION", "initial")
        context.auth.observe_http(200, "VENDOR_HIKVISION", authenticated=True)
        assert context.auth.begin(
            "ONVIF", "missing_identity", missing_identity_fields=("serial",)
        )
        try:
            ControlledOnvifClient(context, authenticated=True).devicemgmt().GetCapabilities(
                Category="All"
            )
            assert context.auth.allows_authenticated_request("ONVIF")

            result = SnapshotService(
                rtsp_provider=RtspSpy(),
                image_processor=lambda _raw: (b"thumbnail", 320, 180),
            ).capture_one(
                target,
                manufacturer="Hikvision",
                model="DS-KD8003",
                context=context,
            )

            assert result.failure is SnapshotFailure.TLS_ERROR
            assert [(attempt.method, attempt.failure) for attempt in result.attempts] == [
                (SnapshotMethod.ONVIF_HTTP, SnapshotFailure.TLS_ERROR)
            ]
            assert context.tls_failed
            assert bad_snapshot_requests == []
            assert vendor_requests == []
        finally:
            assert context.close()
