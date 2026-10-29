"""Additional offline evidence for Panasonic skip and D03 interruption boundaries."""
from __future__ import annotations

from dataclasses import fields
from pathlib import Path

import openpyxl
import pytest
import requests
from requests.adapters import HTTPAdapter

from cam_scanner.application.collection_context import CollectionContext
from cam_scanner.application.inventory_service import InventoryService
from cam_scanner.application.multi_workflow import MultiWorkflow
from cam_scanner.application.single_workflow import SingleWorkflow
from cam_scanner.cameras.registry import build_default_registry
from cam_scanner.cameras.strategy import VendorFirstCollector
from cam_scanner.cameras.transport import TargetTransportPolicy
from cam_scanner.config.models import EffectiveConfig, TransportProfile
from cam_scanner.domain.enums import CollectionStatus, MainMenuChoice
from cam_scanner.domain.models import (
    CameraResult,
    CameraTarget,
    CollectionAttempt,
    CollectionReport,
    parse_model_hint,
)
from cam_scanner.snapshot import rtsp as rtsp_module
from cam_scanner.snapshot.rtsp import RtspFrameProvider
from cam_scanner.snapshot.service import SnapshotCaptureError, SnapshotFailure, SnapshotService
from cam_scanner.terminal.ui import TerminalUI
from tests import test_astra_acl_operation_matrix as acl_fixture
from tests.test_astra_gpn_joint_integration import (
    _CapturingInventory,
    _HttpsFixture,
    _inventory,
    _MultiUI,
    _NoMedia,
    _Progress,
    _trust_fixture_ca,
)

_MESSAGE = "Compatibilidade Panasonic não comprovada; fonte ignorada (CGI não consultada)."


def _skip(report: CollectionReport) -> CollectionAttempt:
    attempts = [a for a in report.attempts if a.operation == "GET /cgi-bin/getinfo"]
    assert len(attempts) == 1
    attempt = attempts[0]
    assert (attempt.stage, attempt.outcome, attempt.error_category) == (
        "VALIDATE", "SKIPPED", "ENDPOINT_UNSUPPORTED",
    )
    assert attempt.http_status is None and attempt.protocol is None
    assert attempt.authentication_mechanism is None
    return attempt


@pytest.mark.parametrize("model", ["", "unknown-model", "WV-S2531", "WV-S2531M"])
def test_single_ineligible_panasonic_shows_skip_without_network(monkeypatch, model):
    outputs = []
    sends = []
    answers = iter(("127.0.0.1", "synthetic-user", "Panasonic", model, "3"))

    class CapturingUI(TerminalUI):
        report = None

        def render_report(self, report, **kwargs):
            self.report = report
            super().render_report(report, **kwargs)

    def forbid_send(*_args, **_kwargs):
        sends.append("send")
        raise AssertionError("inelegible source must not send network requests")

    monkeypatch.setattr(requests.Session, "send", forbid_send)
    ui = CapturingUI(
        input_fn=lambda _prompt: next(answers),
        password_fn=lambda _prompt: "synthetic-password",
        output_fn=outputs.append,
    )
    assert SingleWorkflow(ui, _inventory(443), _NoMedia()).run() is MainMenuChoice.EXIT
    assert ui.report is not None
    _skip(ui.report)
    assert sends == []
    assert _MESSAGE in "\n".join(outputs)
    result = ui.report.to_camera_result("127.0.0.1")
    assert (result.manufacturer, result.model, result.serial) == (None, None, None)
    assert len(fields(CameraResult)) == 14


def test_http_404_is_not_rendered_as_unproven_panasonic_compatibility():
    outputs = []
    ui = TerminalUI(output_fn=outputs.append)
    report = CollectionReport().with_attempt(CollectionAttempt(
        stage="HTTP", operation="GET /cgi-bin/getinfo", outcome="SKIPPED",
        error_category="ENDPOINT_UNSUPPORTED", protocol="HTTPS", http_status=404,
    ))
    ui.render_report(report, ip="127.0.0.1", duration=0.0)
    rendered = "\n".join(outputs)
    assert _MESSAGE not in rendered
    assert "HTTP 404" in rendered


def test_multi_ineligible_models_project_skip_and_never_send(monkeypatch, tmp_path: Path):
    source, destination = tmp_path / "input.xlsx", tmp_path / "output.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(("IP", "USERNAME", "PASSWORD", "FABRICANTE", "MODELO"))
    for index, model in enumerate((None, "unknown-model", "WV-S2531")):
        workbook.active.append(
            ("127.0.0.1", f"user-{index}", f"password-{index}", "Panasonic", model)
        )
    workbook.save(source)
    workbook.close()

    sends = []

    def forbid_send(*_args, **_kwargs):
        sends.append("send")
        raise AssertionError("inelegible source must not send network requests")

    monkeypatch.setattr(requests.Session, "send", forbid_send)
    inventory = _CapturingInventory(_inventory(443))
    workflow = MultiWorkflow(
        _MultiUI(source, destination), inventory, _NoMedia(),
        progress_reporter=_Progress(), max_workers=1,
    )
    workflow.run()
    assert sends == []
    assert len(workflow.completed_rows) == 3 and destination.exists()
    assert len({id(context) for context in inventory.contexts}) == 3
    for row, context in zip(workflow.completed_rows, inventory.contexts, strict=True):
        _skip(row.report)
        assert context.auth.logical_attempts == 0 and context.auth.physical_exchanges == 0
        assert (row.result.manufacturer, row.result.model, row.result.serial) == (None, None, None)
    workbook = openpyxl.load_workbook(destination, data_only=False)
    try:
        diagnostics = list(workbook["DIAGNOSTICOS"].iter_rows(values_only=True))
        assert len(diagnostics[0]) == 9
        skipped = [row for row in diagnostics[1:] if row[3] == "GET /cgi-bin/getinfo"]
        assert len(skipped) == 3
        assert all(
            row[1:6] == (
                "VALIDATE", "N/A", "GET /cgi-bin/getinfo", "SKIPPED", "ENDPOINT_UNSUPPORTED",
            )
            for row in skipped
        )
    finally:
        workbook.close()


def test_multi_complete_identity_does_not_authenticate_for_optional_fields(
    monkeypatch, tmp_path: Path,
):
    """One chain covers profile, Digest, complete identity, XLSX and closed RTSP."""
    monkeypatch.setitem(acl_fixture._ACCOUNTS, "closure-a", ("fixture-pass-a", "Viewer"))
    monkeypatch.setitem(acl_fixture._ACCOUNTS, "closure-b", ("fixture-pass-b", "Viewer"))
    monkeypatch.setattr(acl_fixture, "_ALLOWED_ROLES", frozenset({"Viewer"}))
    source, destination = tmp_path / "input.xlsx", tmp_path / "output.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(("IP", "USERNAME", "PASSWORD", "FABRICANTE", "MODELO"))
    workbook.active.append(("127.0.0.1", "closure-a", "fixture-pass-a", "Hikvision", "Model-F"))
    workbook.active.append(("127.0.0.1", "closure-b", "fixture-pass-b", "Hikvision", "Model-F"))
    workbook.save(source)
    workbook.close()

    with acl_fixture._AclFixture() as fixture:
        profile_id = "closure06-fixture-profile"
        config = EffectiveConfig(transport_profiles=(TransportProfile(
            profile_id, ("127.0.0.1",), "vendor_hikvision", "http", fixture.port,
            "NATIVE_DIGEST_EXISTING", "SYNTHETIC-ACL", True,
        ),))
        inventory = _CapturingInventory(InventoryService(
            VendorFirstCollector(registry=build_default_registry()), config=config,
        ))
        media_calls = []

        class UnsupportedSnapshotProvider:
            def capture(self, *_args, **_kwargs):
                raise SnapshotCaptureError(SnapshotFailure.UNSUPPORTED)

        class ForbiddenLock:
            def acquire(self, *_args, **_kwargs):
                media_calls.append("lock")
                raise AssertionError("Disabled RTSP must stop before the media lock")

            def release(self):
                media_calls.append("release")

        class ForbiddenAv:
            def open(self, *_args, **_kwargs):
                media_calls.append("pyav.open")
                raise AssertionError("Disabled RTSP must stop before PyAV")

        def forbidden_onvif_client(**_kwargs):
            media_calls.append("onvif.media_client")
            raise AssertionError("Disabled RTSP must stop before media ONVIF")

        monkeypatch.setattr(rtsp_module, "_RTSP_CAPTURE_ENABLED", False)
        monkeypatch.setattr(rtsp_module, "_RTSP_CAPTURE_LOCK", ForbiddenLock())
        snapshots = SnapshotService(
            UnsupportedSnapshotProvider(),
            UnsupportedSnapshotProvider(),
            RtspFrameProvider(
                client_factory=forbidden_onvif_client,
                av_module=ForbiddenAv(),
            ),
        )
        workflow = MultiWorkflow(
            _MultiUI(source, destination, profile_id), inventory, snapshots,
            progress_reporter=_Progress(), max_workers=1,
        )
        workflow.run()

    assert destination.exists() and len(workflow.completed_rows) == 2
    assert fixture.requests == [
        ("GET /ISAPI/System/deviceInfo", False), ("GET /ISAPI/System/deviceInfo", True),
        ("GET /ISAPI/System/deviceInfo", False), ("GET /ISAPI/System/deviceInfo", True),
    ]
    assert fixture.authenticated_roles == ["Viewer", "Viewer"]
    assert len({id(context) for context in inventory.contexts}) == 2
    for row, context in zip(workflow.completed_rows, inventory.contexts, strict=True):
        assert context.config.transport_plan.selection_id == profile_id
        assert context.auth.logical_attempts == 1 and context.auth.physical_exchanges == 2
        assert row.report.model_hint.match_key == "model-f"
        assert (row.result.manufacturer, row.result.model, row.result.serial) == (
            "FixtureCam", "Model-F", "Serial-F",
        )
        assert row.result.firmware is None and row.result.mac is None
        assert row.result.hardware_id is None
        assert not context.auth.begin("ONVIF", "missing_identity", missing_identity_fields=())
        assert not context.auth.begin(
            "VENDOR_AXIS", "missing_identity", missing_identity_fields=("serial",),
        )
        assert row.snapshot is not None
        assert row.snapshot.failure is SnapshotFailure.UNSUPPORTED
        assert row.snapshot.attempts == ()
        assert not context.snapshot_eligible("ONVIF")
        assert context.snapshot_eligible("VENDOR_HIKVISION")
    assert media_calls == []
    assert len(fields(CameraResult)) == 14
    workbook = openpyxl.load_workbook(destination, data_only=False)
    try:
        diagnostics = list(workbook["DIAGNOSTICOS"].iter_rows(values_only=True))
        assert len(diagnostics[0]) == 9
        rendered = str([cell for sheet in workbook for row in sheet.iter_rows(values_only=True)
                        for cell in row if cell is not None])
        assert all(value not in rendered for value in (
            "closure-a", "closure-b", "fixture-pass-a", "fixture-pass-b",
        ))
    finally:
        workbook.close()


def test_multi_cancel_after_digest_challenge_does_not_export(monkeypatch, tmp_path: Path):
    source, destination = tmp_path / "input.xlsx", tmp_path / "output.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(("IP", "USERNAME", "PASSWORD", "FABRICANTE", "MODELO"))
    workbook.active.append(
        ("127.0.0.1", "synthetic-user", "synthetic-password", "Panasonic", "WV-S2531L")
    )
    workbook.save(source)
    workbook.close()

    class ContextInventory(_CapturingInventory):
        def __init__(self, service):
            super().__init__(service)
            self.created = []

        def create_context(self, *args, **kwargs):
            context = self.service.create_context(*args, **kwargs)
            self.created.append(context)
            return context

    with _HttpsFixture("digest") as server:
        _trust_fixture_ca(monkeypatch, server.port)
        inventory = ContextInventory(_inventory(server.port))
        trusted_send = HTTPAdapter.send

        def cancel_after_challenge(adapter, request, **kwargs):
            response = trusted_send(adapter, request, **kwargs)
            if len(server.requests) == 1:
                inventory.created[0].cancel_event.set()
            return response

        monkeypatch.setattr(HTTPAdapter, "send", cancel_after_challenge)
        ui = _MultiUI(source, destination)
        workflow = MultiWorkflow(
            ui, inventory, _NoMedia(), progress_reporter=_Progress(), max_workers=1,
        )
        workflow.run()
        assert len(server.requests) == 1
        assert server.requests[0]["path"] == "/cgi-bin/getinfo?FILE=1"
        assert inventory.created[0].auth.logical_attempts == 1
        assert inventory.created[0].auth.physical_exchanges == 1
        assert not destination.exists()
        assert any("cancel" in message.casefold() for message in ui.outputs)


@pytest.mark.parametrize("interruption", ["cancel", "deadline"])
def test_digest_challenge_interruption_blocks_replay(monkeypatch, interruption):
    """A real Requests Digest replay is stopped at the next guarded send."""
    with _HttpsFixture("digest") as server:
        _trust_fixture_ca(monkeypatch, server.port)
        service = _inventory(server.port)
        target = CameraTarget("127.0.0.1", "synthetic-user", "synthetic-password")
        config, blocked, _metadata = service.prepare_transport_for_row(
            target.ip, "Panasonic", "HTTPS_DEFAULT", model_hint="WV-S2531L",
        )
        assert not blocked
        clock = [100.0]
        context = CollectionContext(
            target, config, transport_factory=TargetTransportPolicy,
            clock=lambda: clock[0], manufacturer_hint="Panasonic",
            model_hint=parse_model_hint("WV-S2531L"),
        )
        trusted_send = HTTPAdapter.send

        def interrupt_after_challenge(adapter, request, **kwargs):
            response = trusted_send(adapter, request, **kwargs)
            if len(server.requests) == 1:
                if interruption == "cancel":
                    context.cancel_event.set()
                else:
                    clock[0] = context.budget.deadline + 1.0
            return response

        monkeypatch.setattr(HTTPAdapter, "send", interrupt_after_challenge)
        try:
            report, returned_context = service.collect_execution(target, context=context)
            assert returned_context is context
            assert len(server.requests) == 1
            assert server.requests[0]["path"] == "/cgi-bin/getinfo?FILE=1"
            assert context.auth.logical_attempts == 1
            assert context.auth.physical_exchanges == 1
            assert not context.auth.begin(
                "VENDOR_AXIS", "missing_identity", missing_identity_fields=("serial",),
            )
            assert report.to_camera_result(target.ip).serial is None
            if interruption == "cancel":
                assert report.status is CollectionStatus.CANCELLED
            else:
                assert any(a.error_category == "DEADLINE_EXCEEDED" for a in report.attempts)
        finally:
            context.close()
