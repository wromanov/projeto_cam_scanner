"""Focused offline evidence for the remaining Astra acceptance criteria."""
from __future__ import annotations

import hashlib
import hmac
from dataclasses import fields
from pathlib import Path

import openpyxl
import pytest
import requests
from requests.utils import parse_dict_header

from cam_scanner.application.inventory_service import InventoryService
from cam_scanner.application.multi_workflow import MultiWorkflow
from cam_scanner.application.single_workflow import SingleWorkflow
from cam_scanner.cameras.preauth import PreAuthEvidence
from cam_scanner.cameras.registry import build_default_registry
from cam_scanner.cameras.strategy import VendorFirstCollector
from cam_scanner.domain.enums import MainMenuChoice
from cam_scanner.domain.models import CameraResult, CameraTarget, CollectionReport
from cam_scanner.snapshot.service import SnapshotFailure, SnapshotService
from cam_scanner.terminal.ui import TerminalUI
from tests.test_astra_gpn_joint_integration import (
    _CapturingInventory,
    _HttpsFixture,
    _inventory,
    _MultiUI,
    _NoMedia,
    _Progress,
    _SingleUI,
    _trust_fixture_ca,
)


def test_strong_conflict_in_single_never_selects_a_vendor_or_sends_cgi(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Two observed families conflict before vendor selection, in either order."""
    for families in (("Panasonic", "Axis"), ("Axis", "Panasonic")):
        class Discovery:
            def discover(self, _target, *, context=None, _families=families):
                return PreAuthEvidence(
                    capabilities=[{"manufacturer": family} for family in _families]
                )

        class GenericOnvif:
            calls = 0

            def collect_report(self, _target, *, context=None):
                self.calls += 1
                return CollectionReport()

        class RecordingRegistry:
            def __init__(self):
                self.registry = build_default_registry()
                self.lookups = []

            def get(self, manufacturer):
                self.lookups.append(manufacturer)
                return self.registry.get(manufacturer)

        class UI:
            report = None

            def prompt_camera_target(self):
                return CameraTarget("127.0.0.1", "fixture-user", "fixture-password")

            def prompt_manufacturer_hint(self):
                return None

            def prompt_model_hint(self):
                return "WV-S2531L"

            def render_report(self, report, **_kwargs):
                self.report = report

            def post_query_menu(self):
                return MainMenuChoice.EXIT

        sends = []

        def forbid_send(_session, request, _sends=sends, **_kwargs):
            _sends.append(request.url)
            raise AssertionError("Unexpected network request in conflict fixture")

        monkeypatch.setattr(requests.Session, "send", forbid_send)
        registry = RecordingRegistry()
        generic = GenericOnvif()
        ui = UI()
        service = InventoryService(VendorFirstCollector(
            registry=registry, discoverer=Discovery(), onvif_collector=generic,
        ))
        assert SingleWorkflow(ui, service).run() is MainMenuChoice.EXIT
        assert ui.report is not None
        resolution = ui.report.manufacturer_resolution
        assert resolution is not None
        assert resolution.selected is None and resolution.confidence == "CONFLICTING"
        assert set(resolution.conflicts) == set(families)
        assert ui.report.to_camera_result("127.0.0.1").manufacturer is None
        assert registry.lookups == []
        assert sends == []
        skipped = [a for a in ui.report.attempts if a.operation == "GET /cgi-bin/getinfo"]
        assert len(skipped) == 1
        assert (skipped[0].stage, skipped[0].outcome, skipped[0].error_category) == (
            "VALIDATE", "SKIPPED", "ENDPOINT_UNSUPPORTED",
        )
        assert skipped[0].http_status is None and skipped[0].protocol is None
        assert generic.calls == 1  # Controlled ONVIF fallback; no vendor family is selected.


def test_single_legacy_blank_valid_and_invalid_model_inputs_preserve_target() -> None:
    """Real SINGLE service/context keep hints separate from observed identity."""
    target = CameraTarget("127.0.0.1", "fixture-user", "fixture-password")
    contexts = []
    targets = []

    class Collector:
        def collect_report(self, received, *, context=None):
            targets.append(received)
            contexts.append(context)
            return CollectionReport()

    class LegacyUI:
        report = None

        def prompt_camera_target(self):
            return target

        def prompt_manufacturer_hint(self):
            return None

        def render_report(self, report, **_kwargs):
            self.report = report

        def post_query_menu(self):
            return MainMenuChoice.EXIT

    class NoHintUI(LegacyUI):
        def prompt_model_hint(self):
            return None

    class CapturingTerminal(TerminalUI):
        report = None

        def render_report(self, report, **kwargs):
            self.report = report
            super().render_report(report, **kwargs)

    def terminal(model):
        answers = iter((target.ip, target.username, "", model, "3"))
        return CapturingTerminal(
            input_fn=lambda _prompt: next(answers),
            password_fn=lambda _prompt: target.password,
            output_fn=lambda _text: None,
        )

    uis = (LegacyUI(), NoHintUI(), terminal(""), terminal("WV-S2531L"), terminal("bad\x01model"))
    for ui in uis:
        assert SingleWorkflow(ui, InventoryService(Collector())).run() is MainMenuChoice.EXIT
        assert ui.report is not None
        result = ui.report.to_camera_result(target.ip)
        assert (result.manufacturer, result.model, result.serial) == (None, None, None)
        assert len(fields(CameraResult)) == 14

    assert all(received == target for received in targets)
    assert len({id(context) for context in contexts}) == len(uis)
    assert all(context.target == target for context in contexts)
    assert [context.model_hint.normalized if context.model_hint else None for context in contexts] == [
        None, None, None, "WV-S2531L", None,
    ]
    assert any(attempt.operation == "validate_model_hint" for attempt in uis[-1].report.attempts)


@pytest.mark.parametrize(
    ("mode", "policy", "credential", "firmware", "expected_status"),
    [
        ("single", "allow", "valid", "9.99", 200),
        ("single", "allow", "valid", None, 200),
        ("single", "allow", "invalid", None, 401),
        ("single", "deny", "valid", None, 403),
        ("single", "missing_api", "valid", None, 404),
        ("multi", "allow", "valid", "3.14", 200),
        ("multi", "allow", "invalid", None, 401),
        ("multi", "deny", "valid", None, 403),
        ("multi", "missing_api", "valid", None, 404),
    ],
)
def test_panasonic_digest_fixture_validates_account_and_separates_api_outcomes(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    mode: str,
    policy: str,
    credential: str,
    firmware: str | None,
    expected_status: int,
) -> None:
    """Synthetic loopback ACL validates Digest; it says nothing about real camera ACLs."""
    accounts = {
        "synthetic-allow-05": ("synthetic-allow-password-05", True),
        "synthetic-deny-05": ("synthetic-deny-password-05", False),
    }
    username = "synthetic-deny-05" if policy == "deny" else "synthetic-allow-05"
    correct_password, account_allowed = accounts[username]
    password = correct_password if credential == "valid" else "incorrect-fixture-password"
    realm, nonce = "fixture-realm-05", "fixture-nonce-05"
    server = _HttpsFixture("partial")
    authenticated_accounts = []
    statuses = []

    def reply(handler, status, *, challenge=False, body=b""):
        statuses.append(status)
        handler.send_response(status)
        if challenge:
            handler.send_header(
                "WWW-Authenticate",
                f'Digest realm="{realm}", nonce="{nonce}", qop="auth", algorithm=MD5',
            )
        handler.send_header("Content-Length", str(len(body)))
        handler.end_headers()
        if body:
            handler.wfile.write(body)

    def cgi(handler):
        authorization = handler.headers.get("Authorization")
        if not authorization:
            reply(handler, 401, challenge=True)
            return
        if not authorization.startswith("Digest "):
            reply(handler, 401)
            return
        digest = parse_dict_header(authorization[7:])
        uri = digest.get("uri", "")
        nc, cnonce, qop = digest.get("nc", ""), digest.get("cnonce", ""), digest.get("qop", "")
        valid_fields = (
            digest.get("username") == username
            and digest.get("realm") == realm
            and digest.get("nonce") == nonce
            and uri == "/cgi-bin/getinfo?FILE=1"
            and qop == "auth"
            and bool(nc and cnonce)
        )
        ha1 = hashlib.md5(
            f"{username}:{realm}:{correct_password}".encode(), usedforsecurity=False,
        ).hexdigest()
        ha2 = hashlib.md5(f"GET:{uri}".encode(), usedforsecurity=False).hexdigest()
        expected = hashlib.md5(
            f"{ha1}:{nonce}:{nc}:{cnonce}:{qop}:{ha2}".encode(), usedforsecurity=False,
        ).hexdigest()
        if not valid_fields or not hmac.compare_digest(digest.get("response", ""), expected):
            reply(handler, 401)
            return
        authenticated_accounts.append(username)
        if not account_allowed:
            reply(handler, 403)
        elif policy == "missing_api":
            reply(handler, 404)
        else:
            body = b"NAME=WV-S2531L\nSERIAL=SERIAL-FIXTURE\n"
            if firmware is not None:
                body += f"VERSION={firmware}\n".encode()
            reply(handler, 200, body=body)

    monkeypatch.setattr(server.httpd.RequestHandlerClass, "_cgi", cgi)
    calls = []

    class BlockedProvider:
        def capture(self, *_args, **_kwargs):
            calls.append("image-provider")
            raise AssertionError("Image provider must remain blocked after refusal")

    use_snapshot_guard = policy == "deny" or (mode == "multi" and policy == "allow")
    snapshot_service = (
        SnapshotService(BlockedProvider(), BlockedProvider(), BlockedProvider())
        if use_snapshot_guard else _NoMedia()
    )
    if mode == "multi" and policy == "allow":
        def deny_onvif(handler):
            handler.rfile.read(int(handler.headers.get("Content-Length", "0")))
            server.requests.append({"method": "POST", "path": handler.path, "scheme": "https"})
            handler.send_response(403)
            handler.send_header("Content-Length", "0")
            handler.end_headers()

        monkeypatch.setattr(server.httpd.RequestHandlerClass, "do_POST", deny_onvif)

    class UI(_SingleUI):
        def prompt_camera_target(self):
            return CameraTarget("127.0.0.1", username, password)

        def render_report(self, report, **kwargs):
            self.snapshot = kwargs.get("snapshot")
            super().render_report(report, **kwargs)

    with server:
        _trust_fixture_ca(monkeypatch, server.port)
        inventory = _CapturingInventory(_inventory(server.port))
        if mode == "single":
            ui = UI("WV-S2531L")
            assert SingleWorkflow(ui, inventory, snapshot_service).run() is MainMenuChoice.EXIT
            report = ui.report
            snapshot = ui.snapshot
            outputs = ui.outputs
        else:
            source, destination = tmp_path / "input.xlsx", tmp_path / "output.xlsx"
            workbook = openpyxl.Workbook()
            workbook.active.append(("IP", "USERNAME", "PASSWORD", "FABRICANTE", "MODELO"))
            workbook.active.append(("127.0.0.1", username, password, "Panasonic", "WV-S2531L"))
            workbook.save(source)
            workbook.close()
            ui = _MultiUI(source, destination)
            workflow = MultiWorkflow(
                ui, inventory, snapshot_service, progress_reporter=_Progress(), max_workers=1,
            )
            workflow.run()
            assert len(workflow.completed_rows) == 1
            report = workflow.completed_rows[0].report
            snapshot = workflow.completed_rows[0].snapshot
            assert destination.exists()
            workbook = openpyxl.load_workbook(destination, data_only=False)
            try:
                outputs = [cell for sheet in workbook for row in sheet.iter_rows(values_only=True)
                           for cell in row if cell is not None]
                assert len(next(iter(workbook["DIAGNOSTICOS"].iter_rows(values_only=True)))) == 9
            finally:
                workbook.close()

        assert report is not None
        context = inventory.contexts[0]
        assert context.config.transport_plan.selection_id == "HTTPS_DEFAULT"
        assert context.panasonic_profile == "i-pro-command-interface-v1.17:WV-S2531L"
        cgi_attempts = [a for a in report.attempts if a.operation == "GET /cgi-bin/getinfo"]
        assert len(cgi_attempts) == 1 and cgi_attempts[0].http_status == expected_status
        assert cgi_attempts[0].protocol == "HTTPS"
        assert statuses == [401, expected_status]
        assert len([request for request in server.requests
                    if request["path"] == "/cgi-bin/getinfo?FILE=1"]) == 2
        assert context.auth.logical_attempts <= 2
        assert context.auth.physical_exchanges == len(server.requests)
        assert authenticated_accounts == ([username] if credential == "valid" else [])
        result = report.to_camera_result("127.0.0.1")
        if expected_status == 200:
            assert (result.manufacturer, result.model, result.serial, result.firmware) == (
                None, "WV-S2531L", "SERIAL-FIXTURE", firmware,
            )
            assert any(e.field == "serial" and e.source == "VENDOR" and e.protocol == "HTTPS"
                       for e in report.evidence)
        else:
            assert (result.manufacturer, result.model, result.serial) == (None, None, None)
        if expected_status in (401, 403):
            assert len(server.requests) == 2
            assert not context.auth.begin("ONVIF", "missing_identity", missing_identity_fields=("serial",))
        before = (context.auth.logical_attempts, context.auth.physical_exchanges)
        assert not context.auth.begin(
            "VENDOR_AXIS", "missing_identity", missing_identity_fields=("manufacturer",),
        )
        assert (context.auth.logical_attempts, context.auth.physical_exchanges) == before
        if use_snapshot_guard:
            expected_state = (
                context.auth.AUTH_REJECTED_FINAL
                if expected_status == 401 else context.auth.AUTHORIZATION_DENIED
            )
            assert context.auth.state == expected_state
            assert snapshot is not None and snapshot.failure is SnapshotFailure.AUTH_REJECTED
            assert calls == []
        assert username not in repr(report) and password not in repr(report)
        assert username not in str(outputs) and password not in str(outputs)
