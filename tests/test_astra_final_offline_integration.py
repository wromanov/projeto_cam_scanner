"""Additional Astra cross-component evidence using local synthetic endpoints."""
from __future__ import annotations

import threading
from concurrent.futures import ThreadPoolExecutor
from dataclasses import fields
from pathlib import Path

import openpyxl
import pytest
import requests
from requests.adapters import HTTPAdapter

from cam_scanner.application import multi_workflow as multi_module
from cam_scanner.application.collection_context import CollectionTLSFailure
from cam_scanner.application.inventory_service import InventoryService
from cam_scanner.application.multi_workflow import MultiWorkflow
from cam_scanner.application.single_workflow import SingleWorkflow
from cam_scanner.cameras.registry import build_default_registry
from cam_scanner.cameras.strategy import VendorFirstCollector
from cam_scanner.config.models import EffectiveConfig, TransportProfile
from cam_scanner.domain.enums import CollectionStatus, MainMenuChoice
from cam_scanner.domain.models import (
    CameraResult,
    CollectionAttempt,
    CollectionReport,
    EvidenceRecord,
)
from cam_scanner.snapshot.service import SnapshotFailure, SnapshotService
from cam_scanner.terminal.ui import TerminalUI
from tests.test_astra_gpn_joint_integration import (
    _PASSWORD,
    _USERNAME,
    _CapturingInventory,
    _HttpsFixture,
    _inventory,
    _MultiUI,
    _NoMedia,
    _Progress,
    _SingleUI,
    _trust_fixture_ca,
)
from tests.test_astra_p0_axis_gap_closure import _SUPPORTED, _axis_fixture, _Reply


class _BlockedProvider:
    def __init__(self, calls: list[str], name: str) -> None:
        self.calls, self.name = calls, name

    def capture(self, *_args, **_kwargs):
        self.calls.append(self.name)
        raise AssertionError("A blocked snapshot provider was entered.")


class _SnapshotProbe(SnapshotService):
    def __init__(self, calls: list[str]) -> None:
        super().__init__(
            _BlockedProvider(calls, "onvif"),
            _BlockedProvider(calls, "vendor"),
            _BlockedProvider(calls, "rtsp"),
        )
        self.contexts = []

    def capture_one(self, target, manufacturer=None, model=None, *, context=None):
        self.contexts.append(context)
        return super().capture_one(target, manufacturer, model, context=context)


class _SingleCapture(_SingleUI):
    def render_report(self, report, **kwargs) -> None:
        self.snapshot = kwargs["snapshot"]
        super().render_report(report, **kwargs)


@pytest.mark.parametrize("mode", ["single", "multi"])
@pytest.mark.parametrize("failure", ["401", "403", "tls"])
def test_panasonic_partial_then_onvif_refusal_or_tls_blocks_snapshot_in_same_context(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path, mode: str, failure: str,
) -> None:
    """Observe partial evidence, later failure, ledger and snapshot as one chain."""
    server = _HttpsFixture("partial")

    def denied_post(handler):
        handler.rfile.read(int(handler.headers.get("Content-Length", "0")))
        server.requests.append({"method": "POST", "path": handler.path, "scheme": "https"})
        handler.send_response(int(failure))
        handler.send_header("Content-Length", "0")
        handler.end_headers()

    if failure != "tls":
        monkeypatch.setattr(server.httpd.RequestHandlerClass, "do_POST", denied_post)
    calls: list[str] = []
    snapshots = _SnapshotProbe(calls)
    with server:
        _trust_fixture_ca(monkeypatch, server.port)
        if failure == "tls":
            send = HTTPAdapter.send

            def fail_on_onvif(adapter, request, **kwargs):
                if "/onvif/device_service" in request.url:
                    assert kwargs.get("verify") is True
                    raise requests.exceptions.SSLError("synthetic TLS failure")
                return send(adapter, request, **kwargs)

            monkeypatch.setattr(HTTPAdapter, "send", fail_on_onvif)
        inventory = _CapturingInventory(_inventory(server.port))
        if mode == "single":
            ui = _SingleCapture("WV-S2531L")
            SingleWorkflow(ui, inventory, snapshots).run()
            report, snapshot = ui.report, ui.snapshot
        else:
            source, destination = tmp_path / "input.xlsx", tmp_path / "output.xlsx"
            workbook = openpyxl.Workbook()
            workbook.active.append(("IP", "USERNAME", "PASSWORD", "FABRICANTE", "MODELO"))
            workbook.active.append(("127.0.0.1", _USERNAME, _PASSWORD, "Panasonic", "WV-S2531L"))
            workbook.save(source)
            workbook.close()
            workflow = MultiWorkflow(
                _MultiUI(source, destination), inventory, snapshots,
                progress_reporter=_Progress(), max_workers=1,
            )
            workflow.run()
            assert len(workflow.completed_rows) == 1
            row = workflow.completed_rows[0]
            report, snapshot = row.report, row.snapshot
            output = openpyxl.load_workbook(destination, data_only=False)
            try:
                assert len(next(iter(output["DIAGNOSTICOS"].iter_rows(values_only=True)))) == 9
                assert output["INVENTARIO"]["K2"].value == "PARTIAL_SUCCESS"
            finally:
                output.close()

    assert report is not None and snapshot is not None
    assert len(inventory.contexts) == 1
    context = inventory.contexts[0]
    result = report.to_camera_result("127.0.0.1")
    assert result.status is CollectionStatus.PARTIAL_SUCCESS
    assert (result.manufacturer, result.model, result.serial, result.firmware) == (
        None, "WV-S2531L", "SERIAL-FIXTURE", "9.99",
    )
    assert any(item.source == "VENDOR" and item.field == "serial" for item in report.evidence)
    assert context.auth.logical_attempts <= 2
    assert context.auth.physical_exchanges == 2
    assert not calls
    if failure == "tls":
        assert context.tls_failed
        assert snapshot.failure is SnapshotFailure.TLS_ERROR
        with pytest.raises(CollectionTLSFailure):
            context.check("ONVIF")
        assert [item["method"] for item in server.requests] == ["GET"]
        assert snapshots.contexts == []  # workflow blocks before invoking the snapshot service
    else:
        expected = context.auth.AUTH_REJECTED_FINAL if failure == "401" else context.auth.AUTHORIZATION_DENIED
        assert context.auth.state == expected
        assert not context.auth.begin("ONVIF", "initial")
        assert snapshot.failure is SnapshotFailure.AUTH_REJECTED
        assert snapshots.contexts == [context]
        assert [item["method"] for item in server.requests] == ["GET", "POST"]


def test_single_transport_selection_and_frozen_plan_precede_first_request(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A later configuration swap cannot alter the selected production plan."""
    properties = (
        b'{"data":{"propertyList":{"Brand":"Axis",'
        b'"ProdFullName":"AXIS P14","SerialNumber":"FIXTURE-SERIAL"}}}'
    )
    events: list[str] = []
    outputs: list[str] = []
    answers = iter(("127.0.0.1", "safe-fixture-user", "Axis", "2", "", "3"))

    with _axis_fixture((_Reply(200, _SUPPORTED), _Reply(200, properties))) as (server, observed):
        profile = TransportProfile(
            "axis-local", ("127.0.0.1",), "vendor_axis", "http", server.port,
            "PUBLIC_NONE", "approval-local", True,
        )

        class Service(InventoryService):
            selected_config = None
            context = None

            def prepare_transport_for_row(self, *args, **kwargs):
                selected, blocked, metadata = super().prepare_transport_for_row(*args, **kwargs)
                assert not blocked
                assert metadata["authorization_ref"] == profile.authorization_ref
                assert profile.http_risk_accepted is True
                self.selected_config = selected
                events.append("compiled")
                self._config = EffectiveConfig()  # simulate a later replacement of settings
                return selected, blocked, metadata

            def create_context(self, target, **kwargs):
                self.context = super().create_context(target, **kwargs)
                return self.context

        class UI(TerminalUI):
            def prompt_transport_selection(self, options):
                selected = super().prompt_transport_selection(options)
                assert selected == profile.id
                events.append("selected")
                return selected

        service = Service(
            VendorFirstCollector(registry=build_default_registry()),
            config=EffectiveConfig(transport_profiles=(profile,)),
        )
        ui = UI(
            input_fn=lambda _prompt, choices=answers: next(choices),
            password_fn=lambda _prompt: "safe-fixture-password",
            output_fn=outputs.append,
        )
        handler_class = server.httpd.RequestHandlerClass
        original_post = handler_class.do_POST

        def observed_post(handler):
            events.append("network")
            original_post(handler)

        monkeypatch.setattr(handler_class, "do_POST", observed_post)
        SingleWorkflow(ui, service).run()

    assert events == ["selected", "compiled", "network", "network"]
    assert any("HTTP não oferece confidencialidade" in line for line in outputs)
    assert service.selected_config is not None and service.context is not None
    plan = service.selected_config.transport_plan
    assert plan is service.context.config.transport_plan
    assert plan.selection_id == profile.id
    entry = plan.for_target("127.0.0.1", "vendor_axis")
    assert entry.metadata.authorization_ref == profile.authorization_ref
    assert entry.metadata.scheme == "http" and entry.rule.port == server.port
    assert not service._config.transport_profiles
    assert len(observed) == 2
    assert all(item.authorization is None and item.cookie is None for item in observed)
    assert service.context.auth.logical_attempts == 0


def test_single_four_and_multi_three_statuses_conflicts_and_missing_identity(tmp_path: Path) -> None:
    """Real UI/workflows/XLSX project synthetic reports without treating hints as evidence."""
    def evidence(field: str, value: str) -> EvidenceRecord:
        return EvidenceRecord(field=field, value=value, source="VENDOR", protocol="HTTPS")

    success = CollectionReport().with_evidence(
        evidence("manufacturer", "FixtureCorp"), evidence("model", "M1"), evidence("serial", "S1"),
    )
    partial = CollectionReport().with_evidence(
        evidence("model", "M2"), evidence("model", "M3"), evidence("serial", "S2"),
    )
    failed = CollectionReport().with_evidence(
        EvidenceRecord(field="onvif_reachable", value="REACHABLE", source="ONVIF_PREAUTH"),
    ).with_attempt(CollectionAttempt(
        stage="AUTH", operation="collect_report", outcome="FAILED", error_category="AUTH_REJECTED",
    ))
    cancelled = CollectionReport(cancelled=True).with_evidence(evidence("model", "M4"))
    reports = dict(zip(("127.0.0.1", "127.0.0.2", "127.0.0.3", "127.0.0.4"),
                       (success, partial, failed, cancelled), strict=True))
    expected = (
        CollectionStatus.SUCCESS, CollectionStatus.PARTIAL_SUCCESS,
        CollectionStatus.FAILED, CollectionStatus.CANCELLED,
    )

    class Collector:
        def collect_report(self, target, *, context=None):
            return reports[target.ip]

    for ip, expected_status in zip(reports, expected, strict=True):
        output: list[str] = []
        answers = iter((ip, "fixture-user", "", "Declared-Only", "3"))
        ui = TerminalUI(
            input_fn=lambda _prompt, choices=answers: next(choices),
            password_fn=lambda _prompt: "fixture-password", output_fn=output.append,
        )
        inventory = _CapturingInventory(InventoryService(Collector()))
        assert SingleWorkflow(ui, inventory).run() is MainMenuChoice.EXIT
        report = inventory.contexts[0].report
        result = report.to_camera_result(ip)
        assert result.status is expected_status
        assert f"Status: {expected_status.value}" in output
        assert "Modelo: Declared-Only" not in output
        if expected_status is CollectionStatus.PARTIAL_SUCCESS:
            assert "model" in report.conflicts and result.model is None
            assert any("Conflitos preservados: model" in line for line in output)
        if expected_status is CollectionStatus.FAILED:
            assert result.manufacturer is result.model is result.serial is None
        if expected_status is CollectionStatus.CANCELLED:
            assert any("cancelada" in line.casefold() for line in output)

    source, destination = tmp_path / "statuses-input.xlsx", tmp_path / "statuses-output.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(("IP", "USERNAME", "PASSWORD", "MODELO"))
    for ip in tuple(reports)[:3]:
        workbook.active.append((ip, "fixture-user", "fixture-password", "Declared-Only"))
    workbook.save(source)
    workbook.close()
    multi_ui = _MultiUI(source, destination)
    workflow = MultiWorkflow(
        multi_ui, InventoryService(Collector()), _NoMedia(),
        progress_reporter=_Progress(), max_workers=1,
    )
    workflow.run()
    assert tuple(row.result.status for row in workflow.completed_rows) == expected[:3]
    assert destination.exists(), multi_ui.outputs
    assert len(fields(CameraResult)) == 14
    output = openpyxl.load_workbook(destination, data_only=False)
    try:
        rows = list(output["INVENTARIO"].iter_rows(min_row=2, values_only=True))
        assert [row[10] for row in rows] == [status.value for status in expected[:3]]
        assert (rows[0][3], rows[0][5], rows[0][6]) == ("FixtureCorp", "M1", "S1")
        assert (rows[1][3], rows[1][5], rows[1][6]) == ("N/A", "N/A", "S2")
        assert (rows[2][3], rows[2][5], rows[2][6]) == ("N/A", "N/A", "N/A")
        assert len(next(iter(output["DIAGNOSTICOS"].iter_rows(values_only=True)))) == 9
    finally:
        output.close()


def test_multi_cancelled_row_returns_without_spin_or_export(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path,
) -> None:
    """Force cancellation before input exhaustion and bound failure time."""
    source, destination = tmp_path / "cancel-input.xlsx", tmp_path / "cancel-output.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(("IP", "USERNAME", "PASSWORD"))
    workbook.active.append(("127.0.0.1", "fixture-user", "fixture-password"))
    workbook.save(source)
    workbook.close()
    cancel_events: list[threading.Event] = []

    class Collector:
        def collect_report(self, target, *, context=None):
            cancel_events.append(context.cancel_event)
            return CollectionReport(cancelled=True)

    class CompletedBeforeSubmitReturns(ThreadPoolExecutor):
        def submit(self, fn, /, *args, **kwargs):
            future = super().submit(fn, *args, **kwargs)
            future.result(timeout=5)
            return future

    monkeypatch.setattr(multi_module, "ThreadPoolExecutor", CompletedBeforeSubmitReturns)
    ui = _MultiUI(source, destination)
    workflow = MultiWorkflow(
        ui, InventoryService(Collector()), _NoMedia(), progress_reporter=_Progress(), max_workers=1,
    )
    worker = threading.Thread(target=workflow.run, daemon=True)
    worker.start()
    worker.join(timeout=3)
    if worker.is_alive():
        # Allow the intentionally broken implementation to terminate after detection.
        if cancel_events:
            cancel_events[0].clear()
        worker.join(timeout=5)
        pytest.fail("A cancelled batch did not terminate promptly.", pytrace=False)
    assert len(workflow.completed_rows) == 1
    assert workflow.completed_rows[0].result.status is CollectionStatus.CANCELLED
    assert not destination.exists()
    assert any("cancelado antes da exportação" in item for item in ui.outputs)
