"""G-PN-JOINT offline acceptance through the real application path.

All network traffic is confined to a loopback HTTPS fixture. The registry,
collector, Panasonic adapter, workflows, inventory service, and MULTI Excel
reader/writer are production implementations; only the UI boundary is shimmed.
"""
from __future__ import annotations

import ssl
import threading
from dataclasses import fields
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Self
from urllib.parse import urlsplit

import openpyxl
import pytest
from requests.adapters import HTTPAdapter

from cam_scanner.application.inventory_service import InventoryService
from cam_scanner.application.multi_workflow import MultiWorkflow
from cam_scanner.application.single_workflow import SingleWorkflow
from cam_scanner.cameras.registry import build_default_registry
from cam_scanner.cameras.strategy import VendorFirstCollector
from cam_scanner.config.models import EffectiveConfig, TransportRule
from cam_scanner.domain.enums import CollectionStatus, MainMenuChoice
from cam_scanner.domain.models import CameraResult, CameraTarget
from cam_scanner.snapshot.service import SnapshotFailure, SnapshotResult

_PROFILE = "i-pro-command-interface-v1.17:WV-S2531L"
_USERNAME = "joint-user-sentinel"
_PASSWORD = "joint-password-sentinel"


class _HttpsFixture:
    def __init__(self, scenario: str) -> None:
        self.scenario = scenario
        self.requests: list[dict[str, object]] = []
        owner = self

        class Handler(BaseHTTPRequestHandler):
            def do_GET(self) -> None:
                owner.requests.append({
                    "method": "GET", "path": self.path,
                    "scheme": "https", "authorization": self.headers.get("Authorization"),
                })
                if self.path.startswith("/cgi-bin/getinfo"):
                    self._cgi()
                    return
                self.send_response(404)
                self.send_header("Content-Length", "0")
                self.end_headers()

            def do_POST(self) -> None:
                body = self.rfile.read(int(self.headers.get("Content-Length", "0")))
                operation = "GetCapabilities" if b"GetCapabilities" in body else "OtherSOAP"
                owner.requests.append({
                    "method": "POST", "path": self.path, "scheme": "https",
                    "operation": operation,
                })
                self.send_response(404)
                self.send_header("Content-Length", "0")
                self.end_headers()

            def _cgi(self) -> None:
                attempt = sum(1 for req in owner.requests if req["path"] == self.path)
                authorization = self.headers.get("Authorization")
                if owner.scenario == "digest" and not authorization:
                    self.send_response(401)
                    self.send_header(
                        "WWW-Authenticate",
                        'Digest realm="fixture", nonce="synthetic-nonce", '
                        'qop="auth", algorithm=MD5',
                    )
                    self.send_header("Content-Length", "0")
                    self.end_headers()
                    return
                if owner.scenario == "401" and attempt == 1:
                    self.send_response(401)
                    self.send_header(
                        "WWW-Authenticate",
                        'Digest realm="fixture", nonce="synthetic-nonce", '
                        'qop="auth", algorithm=MD5',
                    )
                    self.send_header("Content-Length", "0")
                    self.end_headers()
                    return
                if owner.scenario == "401":
                    self.send_response(401)
                    self.send_header("Content-Length", "0")
                    self.end_headers()
                    return
                if owner.scenario == "403":
                    self.send_response(403)
                    self.send_header("Content-Length", "0")
                    self.end_headers()
                    return
                if owner.scenario == "404":
                    self.send_response(404)
                    self.send_header("Content-Length", "0")
                    self.end_headers()
                    return
                if owner.scenario == "timeout":
                    threading.Event().wait(0.5)
                payload = b"NAME=WV-S2531L\nSERIAL=SERIAL-FIXTURE\nVERSION=9.99\n"
                if owner.scenario == "firmware_unknown":
                    payload = b"NAME=WV-S2531L\nSERIAL=SERIAL-FIXTURE\n"
                if owner.scenario == "partial_model_missing":
                    payload = b"SERIAL=SERIAL-FIXTURE\n"
                if owner.scenario == "partial":
                    payload = b"NAME=WV-S2531L\nSERIAL=SERIAL-FIXTURE\nVERSION=9.99\n"
                self.send_response(200)
                self.send_header("Content-Type", "text/plain")
                self.send_header("Content-Length", str(len(payload)))
                self.end_headers()
                try:
                    self.wfile.write(payload)
                except (BrokenPipeError, ConnectionResetError):
                    pass

            def log_message(self, *_args: object) -> None:
                pass

        self.httpd = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        tls = ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
        tls.load_cert_chain(
            Path(__file__).with_name("tls_test_server_cardb.pem"),
            Path(__file__).with_name("tls_test_server_cardb.key"),
        )
        self.httpd.socket = tls.wrap_socket(self.httpd.socket, server_side=True)
        self.thread = threading.Thread(target=self.httpd.serve_forever, daemon=True)

    @property
    def port(self) -> int:
        return self.httpd.server_port

    def __enter__(self) -> Self:
        self.thread.start()
        return self

    def __exit__(self, *_exc: object) -> None:
        self.httpd.shutdown()
        self.httpd.server_close()
        self.thread.join(timeout=2)


def _trust_fixture_ca(monkeypatch: pytest.MonkeyPatch, port: int) -> None:
    original_send = HTTPAdapter.send
    ca = str(Path(__file__).with_name("tls_test_ca_cardb.pem"))

    def send(adapter, request, **kwargs):
        parsed = urlsplit(request.url)
        if parsed.scheme == "https" and parsed.hostname == "127.0.0.1" and parsed.port == port:
            assert kwargs.get("verify") is True
            kwargs["verify"] = ca
        return original_send(adapter, request, **kwargs)

    monkeypatch.setattr(HTTPAdapter, "send", send)


def _inventory(
    port: int, *, request_timeout: float = 7.0, camera_deadline: float = 45.0,
) -> InventoryService:
    config = EffectiveConfig(
        request_timeout_seconds=request_timeout,
        camera_soft_deadline_seconds=camera_deadline,
        transport_rules=tuple(
            TransportRule("127.0.0.1", component, "https", None, port)
            for component in ("vendor_panasonic", "onvif_device")
        ),
    )
    return InventoryService(
        VendorFirstCollector(registry=build_default_registry()), config=config
    )


class _SingleUI:
    def __init__(self, model: str | None, manufacturer: str = "Panasonic") -> None:
        self.model = model
        self.manufacturer = manufacturer
        self.report = None
        self.outputs: list[str] = []

    def prompt_camera_target(self) -> CameraTarget:
        return CameraTarget("127.0.0.1", _USERNAME, _PASSWORD)

    def prompt_manufacturer_hint(self) -> str:
        return self.manufacturer

    def prompt_model_hint(self) -> str | None:
        return self.model

    def prompt_transport_selection(self, _options) -> str:
        return "HTTPS_DEFAULT"

    def render_report(self, report, **_kwargs) -> None:
        self.report = report
        self.outputs.append(repr(report))

    def post_query_menu(self) -> MainMenuChoice:
        return MainMenuChoice.EXIT


class _MultiUI:
    def __init__(self, source: Path, destination: Path) -> None:
        self.source, self.destination = source, destination
        self.outputs: list[str] = []

    def prompt_multi_input_path(self) -> str:
        return str(self.source)

    def prompt_multi_output_path(self, _default: str) -> str:
        return str(self.destination)

    def prompt_transport_selection(self, _options) -> str:
        return "HTTPS_DEFAULT"

    def show_multi_complete(self, path: Path, total: int, *, failed: int) -> None:
        self.outputs.append(f"{path.name}:{total}:{failed}")

    def show_multi_error(self, message: str) -> None:
        self.outputs.append(message)


class _Progress:
    def start(self, _total: int) -> None: pass
    def advance(self, _ip: str, _status: str) -> None: pass
    def finish(self) -> None: pass


class _NoMedia:
    def capture_one(self, *_args: object, **_kwargs: object) -> SnapshotResult:
        return SnapshotResult(failure=SnapshotFailure.UNSUPPORTED)


def _assert_sentinels_absent(value: object) -> None:
    text = str(value)
    assert _USERNAME not in text
    assert _PASSWORD not in text


def test_exact_model_single_workflow_uses_real_default_panasonic_adapter(monkeypatch) -> None:
    with _HttpsFixture("partial") as server:
        _trust_fixture_ca(monkeypatch, server.port)
        ui = _SingleUI("WV-S2531L")
        assert SingleWorkflow(ui, _inventory(server.port), _NoMedia()).run() is MainMenuChoice.EXIT

    report = ui.report
    assert report is not None
    result = report.to_camera_result("127.0.0.1")
    assert result.status is CollectionStatus.PARTIAL_SUCCESS, [
        (a.operation, a.protocol, a.http_status, a.error_category, dict(a.metadata))
        for a in report.attempts
    ]
    assert (result.manufacturer, result.model, result.serial, result.firmware) == (
        None, "WV-S2531L", "SERIAL-FIXTURE", "9.99",
    )
    assert not any(item.get("operation") == "GetCapabilities" for item in server.requests)
    assert [item["path"] for item in server.requests] == [
        "/cgi-bin/getinfo?FILE=1", "/onvif/device_service",
    ]
    assert server.requests[0]["scheme"] == "https"
    cgi_attempt = next(a for a in report.attempts if a.operation == "GET /cgi-bin/getinfo")
    assert cgi_attempt.http_status == 200
    assert cgi_attempt.error_category is None
    assert len(fields(CameraResult)) == 14
    _assert_sentinels_absent(ui.outputs)
    _assert_sentinels_absent(report)


def test_unknown_firmware_remains_unknown_in_exact_model_response(monkeypatch) -> None:
    with _HttpsFixture("firmware_unknown") as server:
        _trust_fixture_ca(monkeypatch, server.port)
        ui = _SingleUI("WV-S2531L")
        SingleWorkflow(ui, _inventory(server.port), _NoMedia()).run()

    assert ui.report.to_camera_result("127.0.0.1").firmware is None
    cgi_attempt = next(a for a in ui.report.attempts if a.operation == "GET /cgi-bin/getinfo")
    assert cgi_attempt.http_status == 200 and cgi_attempt.error_category is None


def test_model_hint_does_not_fill_unobserved_model_or_manufacturer(monkeypatch) -> None:
    with _HttpsFixture("partial_model_missing") as server:
        _trust_fixture_ca(monkeypatch, server.port)
        ui = _SingleUI("WV-S2531L")
        SingleWorkflow(ui, _inventory(server.port), _NoMedia()).run()

    result = ui.report.to_camera_result("127.0.0.1")
    assert result.model is None and result.manufacturer is None
    assert result.serial == "SERIAL-FIXTURE"
    _assert_sentinels_absent(ui.outputs)


@pytest.mark.parametrize("model", [None, "unknown-model", "WV-S2531", "WV-S2531M"])
def test_ineligible_model_does_not_call_panasonic_cgi_or_capabilities(monkeypatch, model) -> None:
    with _HttpsFixture("partial") as server:
        _trust_fixture_ca(monkeypatch, server.port)
        ui = _SingleUI(model)
        SingleWorkflow(ui, _inventory(server.port), _NoMedia()).run()

    assert not any(item["path"].startswith("/cgi-bin/") for item in server.requests)
    capability_calls = [item for item in server.requests
                        if item.get("operation") == "GetCapabilities"]
    assert not capability_calls, {
        "requests": [(item["method"], item["path"], item["scheme"], item.get("operation"))
                     for item in server.requests],
        "attempts": [
            (a.operation, a.protocol, a.http_status, a.error_category, a.outcome)
            for a in ui.report.attempts
        ] if ui.report is not None else [],
    }
    assert not server.requests
    _assert_sentinels_absent(ui.outputs)


@pytest.mark.parametrize(
    ("scenario", "status", "category"),
    [("404", 404, "ENDPOINT_UNSUPPORTED"), ("401", 401, "AUTH_REJECTED"),
     ("403", 403, "AUTHORIZATION_DENIED")],
)
def test_cgi_http_statuses_are_reported_without_fallback(monkeypatch, scenario, status, category):
    with _HttpsFixture(scenario) as server:
        _trust_fixture_ca(monkeypatch, server.port)
        ui = _SingleUI("WV-S2531L")
        SingleWorkflow(ui, _inventory(server.port), _NoMedia()).run()

    report = ui.report
    assert report is not None
    relevant = [attempt for attempt in report.attempts if attempt.operation == "GET /cgi-bin/getinfo"]
    assert relevant
    assert relevant[-1].http_status == status
    assert relevant[-1].error_category == category
    assert server.requests[0]["path"] == "/cgi-bin/getinfo?FILE=1"
    assert len([item for item in server.requests if str(item["path"]).startswith("/cgi-bin/")]) == (
        2 if scenario == "401" else 1
    )
    _assert_sentinels_absent(ui.outputs)


def test_digest_challenge_then_200_records_both_exchanges_and_effective_https(monkeypatch) -> None:
    with _HttpsFixture("digest") as server:
        _trust_fixture_ca(monkeypatch, server.port)
        ui = _SingleUI("WV-S2531L")
        SingleWorkflow(ui, _inventory(server.port), _NoMedia()).run()

    assert [item["scheme"] for item in server.requests] == ["https", "https", "https"]
    assert [item["path"] for item in server.requests] == [
        "/cgi-bin/getinfo?FILE=1", "/cgi-bin/getinfo?FILE=1", "/onvif/device_service",
    ]
    assert server.requests[0]["authorization"] is None
    assert str(server.requests[1]["authorization"]).startswith("Digest ")
    cgi_attempt = next(a for a in ui.report.attempts if a.operation == "GET /cgi-bin/getinfo")
    assert cgi_attempt.http_status == 200
    assert cgi_attempt.error_category is None
    _assert_sentinels_absent(ui.outputs)


def test_multi_real_excel_path_isolates_credentials_for_same_ip_and_sanitizes_output(
    monkeypatch, tmp_path: Path,
) -> None:
    with _HttpsFixture("digest") as server:
        _trust_fixture_ca(monkeypatch, server.port)
        source, destination = tmp_path / "input.xlsx", tmp_path / "result.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(("IP", "USERNAME", "PASSWORD", "FABRICANTE", "MODELO"))
        sheet.append(("127.0.0.1", "multi-user-a", "multi-secret-a", "Panasonic", "WV-S2531L"))
        sheet.append(("127.0.0.1", "multi-user-b", "multi-secret-b", "Panasonic", "WV-S2531L"))
        workbook.save(source)
        workbook.close()
        ui = _MultiUI(source, destination)
        workflow = MultiWorkflow(
            ui, _inventory(server.port), _NoMedia(), progress_reporter=_Progress(), max_workers=1,
        )
        workflow.run()

    cgi = [item for item in server.requests if str(item["path"]).startswith("/cgi-bin/")]
    assert len(cgi) == 4
    assert cgi[0]["authorization"] is None and cgi[2]["authorization"] is None
    assert 'username="multi-user-a"' in str(cgi[1]["authorization"])
    assert 'username="multi-user-b"' in str(cgi[3]["authorization"])
    assert [row.result.status for row in workflow.completed_rows] == [
        CollectionStatus.PARTIAL_SUCCESS, CollectionStatus.PARTIAL_SUCCESS,
    ]
    assert len(fields(CameraResult)) == 14
    output = openpyxl.load_workbook(destination, data_only=False)
    try:
        values = [cell for row in output.active.iter_rows() for cell in row]
        serialized = "\n".join(str(cell.value) for cell in values if cell.value is not None)
        assert _USERNAME not in serialized and _PASSWORD not in serialized
        assert "multi-user-a" not in serialized and "multi-user-b" not in serialized
        assert "multi-secret-a" not in serialized and "multi-secret-b" not in serialized
        assert sum(1 for row in output.active.iter_rows(min_row=2) if row[0].value is not None) == 2
    finally:
        output.close()
    assert all(item["scheme"] == "https" for item in server.requests)
    _assert_sentinels_absent(ui.outputs)


def test_invalid_tls_is_fail_closed_without_http_fallback() -> None:
    with _HttpsFixture("partial") as server:
        ui = _SingleUI("WV-S2531L")
        SingleWorkflow(ui, _inventory(server.port), _NoMedia()).run()

    report = ui.report
    assert report is not None
    failures = [attempt for attempt in report.attempts if attempt.error_category == "TLS_ERROR"]
    assert failures
    assert all(item["scheme"] == "https" for item in server.requests)
    assert not any(item["path"] == "/onvif/device_service" for item in server.requests)
    _assert_sentinels_absent(ui.outputs)


def test_request_timeout_is_classified_and_preserves_request_sequence(monkeypatch) -> None:
    with _HttpsFixture("timeout") as server:
        _trust_fixture_ca(monkeypatch, server.port)
        ui = _SingleUI("WV-S2531L")
        SingleWorkflow(
            ui, _inventory(server.port, request_timeout=0.1), _NoMedia(),
        ).run()

    report = ui.report
    assert report is not None
    assert any(attempt.error_category in {"READ_TIMEOUT", "CONNECT_TIMEOUT"}
               for attempt in report.attempts)
    assert server.requests[0]["path"] == "/cgi-bin/getinfo?FILE=1"
    assert all(item["scheme"] == "https" for item in server.requests)
    _assert_sentinels_absent(ui.outputs)


def test_expired_deadline_and_cancelled_context_make_no_network_request(monkeypatch) -> None:
    with _HttpsFixture("partial") as server:
        _trust_fixture_ca(monkeypatch, server.port)
        service = _inventory(server.port, camera_deadline=0.05)
        target = CameraTarget("127.0.0.1", _USERNAME, _PASSWORD)
        context = service.create_context(
            target, manufacturer_hint="Panasonic", model_hint="WV-S2531L",
        )
        threading.Event().wait(0.06)
        try:
            report, _ = service.collect_execution(target, context=context)
        finally:
            context.close()
        assert any(attempt.error_category in {"DEADLINE_EXCEEDED", "CONNECT_TIMEOUT"}
                   for attempt in report.attempts)
        assert not server.requests

        cancel_event = threading.Event()
        cancel_event.set()
        cancelled, cancel_context = service.collect_execution(
            target, cancel_event=cancel_event,
            manufacturer_hint="Panasonic", model_hint="WV-S2531L",
        )
        try:
            assert cancelled.status is CollectionStatus.CANCELLED
        finally:
            cancel_context.close()
    assert not server.requests
