"""Consolidated, password-free XLSX output boundary for MULTI."""
from __future__ import annotations

import os
import re
import tempfile
from contextlib import suppress
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import TYPE_CHECKING

from cam_scanner.domain.enums import CollectionStatus
from cam_scanner.domain.models import (
    CameraResult,
    CollectionReport,
    is_safe_xml_text,
    is_valid_evidence_value,
)

if TYPE_CHECKING:
    from cam_scanner.snapshot.service import SnapshotResult

_HEADERS = (
    "LINHA",
    "IP",
    "HOSTNAME",
    "FABRICANTE",
    "FABRICANTE_ORIGINAL",
    "MODELO",
    "NUMERO_SERIE",
    "FIRMWARE",
    "HARDWARE_ID",
    "MAC",
    "STATUS",
    "METODO_COLETA",
    "ERRO_CODIGO",
    "ERRO",
    "DURACAO_SEGUNDOS",
    "SNAPSHOT",
)


class ExcelOutputError(ValueError):
    """Sanitized workbook output error."""


@dataclass(frozen=True, slots=True)
class InventoryOutputRow:
    source_row: int
    result: CameraResult
    snapshot: SnapshotResult | None = None
    report: CollectionReport | None = None


class ExcelWriter:
    """Write ordered inventory and summary sheets without input credentials."""

    def write(self, path: str | Path, rows: list[InventoryOutputRow]) -> Path:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
        except ImportError:
            raise ExcelOutputError("O suporte XLSX não está disponível neste ambiente.") from None
        try:
            from openpyxl.drawing.image import Image as SpreadsheetImage
        except ImportError:
            SpreadsheetImage = None

        destination = Path(path).expanduser()
        if destination.exists():
            raise ExcelOutputError("O arquivo de saída já existe; escolha outro nome.")
        try:
            destination.parent.mkdir(parents=True, exist_ok=True)
        except OSError:
            raise ExcelOutputError("Não foi possível preparar a pasta de saída.") from None

        workbook = Workbook()
        image_streams = []
        temporary_path: Path | None = None
        snapshot_count = 0
        try:
            inventory = workbook.active
            inventory.title = "INVENTARIO"
            inventory.append(_HEADERS)
            for cell in inventory[1]:
                cell.font = Font(bold=True)
            inventory.freeze_panes = "A2"
            inventory.auto_filter.ref = f"A1:P{max(1, len(rows) + 1)}"

            effective_statuses = []
            for row_number, output_row in enumerate(rows, start=2):
                try:
                    result = output_row.result
                    snapshot = output_row.snapshot
                    effective_status = _effective_status(result)
                    has_snapshot = snapshot is not None and snapshot.successful
                    snapshot_label = "Incluído" if has_snapshot else "N/A"
                    values = (
                        _source_row(output_row.source_row),
                        _cell(result.ip),
                        _identity_cell(result.hostname, "hostname"),
                        _identity_cell(result.manufacturer, "manufacturer"),
                        _identity_cell(result.manufacturer_original, "manufacturer_original"),
                        _identity_cell(result.model, "model"),
                        _identity_cell(result.serial, "serial"),
                        _identity_cell(result.firmware, "firmware"),
                        _identity_cell(result.hardware_id, "hardware_id"),
                        _identity_cell(result.mac, "mac"),
                        effective_status.value,
                        result.collection_method.value if result.collection_method else "N/A",
                        result.error_code.value if result.error_code else "N/A",
                        _cell(_inventory_diagnostic_message(output_row)),
                        result.duration,
                        snapshot_label,
                    )
                    inventory.append(values)
                except Exception:  # noqa: BLE001 - isolate malformed legacy rows.
                    inventory.append(
                        (
                            _source_row(output_row.source_row),
                            *("Dado inválido",) * 9,
                            "FAILED",
                            *("Dado inválido",) * 5,
                        )
                    )
                    effective_status = CollectionStatus.FAILED
                    has_snapshot = False
                effective_statuses.append(effective_status)
                if has_snapshot:
                    if SpreadsheetImage is None:
                        inventory.cell(row=row_number, column=16, value="Não incorporado")
                        continue
                    try:
                        stream = BytesIO(snapshot.jpeg_bytes)
                        image_streams.append(stream)
                        image = SpreadsheetImage(stream)
                        image.width = min(image.width, 320)
                        image.height = min(image.height, 180)
                        image.anchor = f"P{row_number}"
                        inventory.add_image(image)
                        inventory.row_dimensions[row_number].height = 105
                        snapshot_count += 1
                    except Exception:  # noqa: BLE001 - one bad image cannot abort the inventory export.
                        inventory.cell(row=row_number, column=16, value="Não incorporado")

            inventory.column_dimensions["A"].width = 10
            inventory.column_dimensions["B"].width = 18
            for column in "CDEFGHIJKLMNO":
                inventory.column_dimensions[column].width = 20
            inventory.column_dimensions["P"].width = 42

            summary = workbook.create_sheet("RESUMO")
            summary.append(("INDICADOR", "QUANTIDADE"))
            for cell in summary[1]:
                cell.font = Font(bold=True)
            counts = {
                "TOTAL": len(effective_statuses),
                "SUCCESS": sum(status.value == "SUCCESS" for status in effective_statuses),
                "PARTIAL_SUCCESS": sum(
                    status.value == "PARTIAL_SUCCESS" for status in effective_statuses
                ),
                "FAILED": sum(status.value == "FAILED" for status in effective_statuses),
                "CANCELLED": sum(status.value == "CANCELLED" for status in effective_statuses),
                "SNAPSHOT_EMBEDDED": snapshot_count,
            }
            for label, count in counts.items():
                summary.append((label, count))
            summary.column_dimensions["A"].width = 26
            summary.column_dimensions["B"].width = 18
            summary.freeze_panes = "A2"

            diagnostics = workbook.create_sheet("DIAGNOSTICOS")
            diagnostics_headers = (
                "LINHA", "ETAPA", "PROTOCOLO", "OPERACAO", "OUTCOME",
                "CATEGORIA", "DURACAO", "CODIGO", "STATUS",
            )
            diagnostics.append(diagnostics_headers)
            for cell in diagnostics[1]:
                cell.font = Font(bold=True)
            diagnostics.freeze_panes = "A2"
            diagnostics.auto_filter.ref = f"A1:I{max(1, sum(len(row.report.attempts) for row in rows if row.report) + 1)}"
            for output_row in sorted(rows, key=lambda item: (
                item.source_row if type(item.source_row) is int and item.source_row >= 0 else 2**63
            )):
                report = output_row.report
                if report is None:
                    continue
                for attempt in report.attempts:
                    diagnostics.append((
                        _source_row(output_row.source_row),
                        _enum_cell(attempt.stage, _DIAGNOSTIC_STAGES),
                        _enum_cell(attempt.protocol, _DIAGNOSTIC_PROTOCOLS),
                        _enum_cell(attempt.operation, _DIAGNOSTIC_OPERATIONS),
                        _enum_cell(attempt.outcome, _DIAGNOSTIC_OUTCOMES),
                        _enum_cell(attempt.error_category, _DIAGNOSTIC_CATEGORIES),
                        _safe_duration(attempt.duration),
                        _safe_diagnostic_code(attempt.protocol_code),
                        _safe_http_status(attempt.http_status),
                    ))
            for column, width in (("A", 10), ("B", 14), ("C", 12), ("D", 48),
                                  ("E", 14), ("F", 28), ("G", 14), ("H", 28), ("I", 12)):
                diagnostics.column_dimensions[column].width = width

            with tempfile.NamedTemporaryFile(
                prefix="cam_scanner_",
                suffix=".xlsx.tmp",
                dir=destination.parent,
                delete=False,
            ) as temporary_file:
                temporary_path = Path(temporary_file.name)
            workbook.save(temporary_path)
            try:
                os.link(temporary_path, destination)
            except FileExistsError:
                raise ExcelOutputError(
                    "O arquivo de saída já existe; escolha outro nome."
                ) from None
            temporary_path.unlink()
            temporary_path = None
            return destination
        except ExcelOutputError:
            raise
        except Exception:  # noqa: BLE001 - library/filesystem details are not user-facing.
            raise ExcelOutputError("Não foi possível gerar o arquivo XLSX de saída.") from None
        finally:
            with suppress(Exception):
                workbook.close()
            for stream in image_streams:
                with suppress(Exception):
                    stream.close()
            if temporary_path is not None:
                with suppress(OSError):
                    temporary_path.unlink(missing_ok=True)


def _cell(value: str | None) -> str:
    if value is None or not value.strip():
        return "N/A"
    if not is_safe_xml_text(value):
        return "Dado inválido"
    # Formula detection runs after validation, including leading whitespace.
    if value.lstrip().startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _identity_cell(value: str | None, field_name: str) -> str:
    if value is None:
        return "N/A"
    if not is_valid_evidence_value(field_name, value):
        return "Dado inválido"
    return _cell(value)


def _source_row(value: object) -> int | str:
    if type(value) is int and value >= 0:
        return value
    return "Dado inválido"


def _inventory_diagnostic_message(row: InventoryOutputRow) -> str | None:
    message = row.result.error_message
    resolution = row.report.manufacturer_resolution if row.report is not None else None
    if resolution is not None and resolution.mismatch:
        warning = "Fabricante declarado diverge do observado."
        return f"{message} {warning}" if message else warning
    return message


_DIAGNOSTIC_STAGES = {"DISCOVERY", "CONNECT", "HTTP", "AUTH", "SOAP", "PARSE", "MERGE", "VALIDATE"}
_DIAGNOSTIC_PROTOCOLS = {"ONVIF", "HTTP", "HTTPS", "RTSP"}
_DIAGNOSTIC_OUTCOMES = {"SUCCESS", "FAILED", "CANCELLED", "SKIPPED"}
_DIAGNOSTIC_CATEGORIES = {
    "CONNECT_TIMEOUT", "READ_TIMEOUT", "CONNECTION_REFUSED", "TLS_ERROR", "HTTP_ERROR",
    "AUTH_REJECTED", "AUTHORIZATION_DENIED", "AUTH_MECHANISM_UNSUPPORTED",
    "ENDPOINT_UNSUPPORTED", "SOAP_FAULT", "INVALID_RESPONSE", "DEADLINE_EXCEEDED",
    "INVALID_INPUT", "UNEXPECTED",
}
_DIAGNOSTIC_OPERATIONS = {
    "GetCapabilities", "GetSystemDateAndTime", "GetDeviceInformation", "create_onvif_client",
    "validate_target", "collect_one", "collect", "collect_result", "collect_report",
    "parse_device_information", "resolve_preauth_manufacturer", "aggregate_collection_report",
    "discover", "preauth_discovery", "vendor_adapter", "ONVIF pre-authentication discovery",
    "GET /ISAPI/System/deviceInfo",
    "POST /axis-cgi/basicdeviceinfo.cgi getAllProperties", "GET /stw-cgi/system.cgi deviceinfo view",
    "GET /cgi-bin/magicBox.cgi getSystemInfo", "GET /cgi-bin/magicBox.cgi getVendor",
    "GET /cgi-bin/getinfo",
}
_DIAGNOSTIC_CODE_PATTERN = re.compile(
    r"^(?:[A-Za-z][A-Za-z0-9]{0,7}:)?(?:Sender|Receiver|MustUnderstand|VersionMismatch|"
    r"NotAuthorized|InvalidArgVal|ActionNotSupported|OperationProhibited|ResourceUnknown)$|"
    r"^(?:0x[0-9A-Fa-f]{1,8}|E[0-9]{1,6})$"
)


def _enum_cell(value: str | None, allowed: set[str]) -> str:
    return value if isinstance(value, str) and value in allowed else "N/A"


def _safe_duration(value: object) -> float | str:
    if type(value) in {int, float} and 0 <= value < float("inf"):
        return float(value)
    return "N/A"


def _safe_http_status(value: object) -> int | str:
    return value if type(value) is int and 100 <= value <= 599 else "N/A"


def _safe_diagnostic_code(value: str | None) -> str:
    if not isinstance(value, str) or not _DIAGNOSTIC_CODE_PATTERN.fullmatch(value):
        return "N/A"
    return _cell(value)


def _effective_status(result: CameraResult):
    if result.status.value != "SUCCESS":
        return result.status
    # Reuse the canonical IDENTITY_V1 semantics for legacy results. Invalid or
    # missing essential identity is exported as partial/failed, never success.
    return CollectionReport.from_camera_result(result).status
