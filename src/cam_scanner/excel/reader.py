"""Streaming XLSX input boundary for bounded MULTI processing."""
from __future__ import annotations

import ipaddress
from collections.abc import Iterator
from dataclasses import dataclass, field
from pathlib import Path

_REQUIRED_HEADERS = ("ip", "username", "password")


class ExcelInputError(ValueError):
    """Sanitized workbook-level input error."""


@dataclass(frozen=True, slots=True)
class ExcelInputRow:
    row_id: int
    ip: str
    username: str = field(repr=False)
    password: str = field(repr=False)
    validation_error: str | None = None

    @property
    def valid(self) -> bool:
        return self.validation_error is None


@dataclass(frozen=True, slots=True)
class ExcelInputBatch:
    total_rows: int
    rows: Iterator[ExcelInputRow] = field(repr=False, compare=False)


class ExcelReader:
    """Read required headers case-insensitively and yield one credential row at a time."""

    def read(self, path: str | Path) -> ExcelInputBatch:
        workbook = self._open(path)
        try:
            worksheet = workbook.active
            header_map = _header_map(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
            total_rows = sum(
                1
                for values in worksheet.iter_rows(min_row=2, values_only=True)
                if _row_has_required_data(values, header_map)
            )
        except ExcelInputError:
            raise
        except Exception:  # noqa: BLE001 - workbook library errors are normalized at the boundary.
            raise ExcelInputError("Não foi possível validar a planilha XLSX de entrada.") from None
        finally:
            workbook.close()
        return ExcelInputBatch(total_rows, self._iter_rows(path, header_map))

    def _iter_rows(self, path: str | Path, expected_headers: dict[str, int]) -> Iterator[ExcelInputRow]:
        workbook = self._open(path)
        try:
            worksheet = workbook.active
            # Re-read the header to detect a file change between counting and processing.
            header_map = _header_map(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
            if header_map != expected_headers:
                raise ExcelInputError("Os cabeçalhos da planilha mudaram durante a leitura.")
            for row_number, values in enumerate(
                worksheet.iter_rows(min_row=2, values_only=True), start=2
            ):
                if not _row_has_required_data(values, header_map):
                    continue
                yield _parse_row(row_number, values, header_map)
        except ExcelInputError:
            raise
        except Exception:  # noqa: BLE001 - malformed worksheet errors are normalized.
            raise ExcelInputError("Não foi possível ler as linhas da planilha XLSX.") from None
        finally:
            workbook.close()

    @staticmethod
    def _open(path: str | Path):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ExcelInputError("O suporte XLSX não está disponível neste ambiente.") from None
        try:
            return load_workbook(Path(path), read_only=True, data_only=True)
        except Exception:  # noqa: BLE001 - filesystem/format details are not user-facing.
            raise ExcelInputError("Não foi possível abrir a planilha XLSX de entrada.") from None


def _header_map(rows) -> dict[str, int]:
    header_row = next(iter(rows), None)
    if header_row is None:
        raise ExcelInputError("A planilha está vazia; informe IP, USERNAME e PASSWORD.")
    normalized = [_cell_text(value).strip().casefold() for value in header_row]
    mapping: dict[str, int] = {}
    for required in _REQUIRED_HEADERS:
        matches = [index for index, value in enumerate(normalized) if value == required]
        if len(matches) != 1:
            raise ExcelInputError("A planilha deve conter exatamente uma coluna IP, USERNAME e PASSWORD.")
        mapping[required] = matches[0]
    return mapping


def _parse_row(row_id: int, values: tuple[object, ...], header_map: dict[str, int]) -> ExcelInputRow:
    ip = _cell_text(_get_cell(values, header_map["ip"])).strip()
    username = _cell_text(_get_cell(values, header_map["username"])).strip()
    password = _cell_text(_get_cell(values, header_map["password"]))
    if not ip or not username or not password.strip():
        return ExcelInputRow(row_id, _safe_ip(ip, username, password), "", "", "missing_required_value")
    try:
        ipaddress.ip_address(ip)
    except ValueError:
        return ExcelInputRow(row_id, _safe_ip(ip, username, password), "", "", "invalid_ip")
    return ExcelInputRow(row_id, ip, username, password)


def _get_cell(values: tuple[object, ...], index: int) -> object:
    return values[index] if index < len(values) else None


def _row_has_required_data(values: tuple[object, ...], header_map: dict[str, int]) -> bool:
    return any(
        _cell_text(_get_cell(values, header_map[header])).strip()
        for header in _REQUIRED_HEADERS
    )


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _safe_ip(value: str, username: str, password: str) -> str:
    folded_value = value.casefold()
    if (
        "://" in folded_value
        or "@" in folded_value
        or folded_value.startswith("www.")
        or any(secret and secret.casefold() in folded_value for secret in (username, password))
    ):
        return "<invalid>"
    return value or "N/A"
